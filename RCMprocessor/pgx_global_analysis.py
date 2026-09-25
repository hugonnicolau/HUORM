from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.worksheet.worksheet import Worksheet

# Caracteres de controlo vindos da conversão, removidos à entrada da folha.
#
# O docling deixa passar caracteres de controlo do PDF para o markdown, e eles
# sobrevivem até aos campos de texto do output. O XML do xlsx não os admite, e
# o openpyxl levanta IllegalCharacterError a meio da escrita — a corrida
# nacional inteira processada e nenhum Excel no fim. Aconteceu com o grupo
# farmacoterapêutico "9.1.2 \x07 Aparelho locomotor...".
#
# A limpeza é feita aqui, num sítio só, e não nas 83 chamadas a ws.append():
# em cada uma delas seria esquecida na primeira que se acrescentasse.
#
# Não altera texto legítimo: o padrão só apanha \x00-\x08, \x0b-\x0c e
# \x0e-\x1f, que nunca aparecem em prosa. Tabulações, mudanças de linha e
# retornos ficam intactos.
_append_original = Worksheet.append


def _append_sem_controlos(self, iterable):
    if isinstance(iterable, (list, tuple)):
        iterable = [
            ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v
            for v in iterable
        ]
    return _append_original(self, iterable)


Worksheet.append = _append_sem_controlos


class PGxGlobalAnalysis:
    """
    Gera análise estatística global a partir de outputs_batch.

    Lógica:
    - Document_Unique_PGx.json:
        entidades únicas, identificadores, substância ativa,
        grupo farmacoterapêutico, ATC e guideline IDs.

    - JSONs individuais das secções:
        ocorrências reais por secção e menções totais reais.

    - Extended_PGx_Analysis.json:
        normalizações detalhadas e informação complementar.

    Identificadores:
    - usa umls_cui como nome oficial.
    - lê mrconso_cui apenas como fallback para compatibilidade com ficheiros antigos.
    """

    PGX_OUTPUTS_DIRNAME = "pgx_outputs"

    UNIQUE_FILENAME = "Document_Unique_PGx.json"
    EXTENDED_FILENAMES = [
        "Extended_PGx_Analysis.json",
        "Resumo_PGx_Extended.json",
    ]

    JSON_OUTPUT = "Analise_Global_PGx.json"
    XLSX_OUTPUT = "Analise_Global_PGx.xlsx"

    SPECIAL_FILES = {
        "Document_Unique_PGx.json",
        "Extended_PGx_Analysis.json",
        "Resumo_PGx_Extended.json",
        "Resumo_Entidades_PGx.json",
        "Analise_Global_PGx.json",
        "Nome_do_Medicamento.json",
        "Grupo_farmacoterapeutico.json",
        "ATC.json",
        "Substancia_ativa.json",
        "Substância_ativa.json",
    }

    def __init__(self, outputs_root: str | Path):
        self.outputs_root = Path(outputs_root).resolve()

        # Populated by _find_doc_dirs. Initialised here so the attribute exists
        # even if the summary is built before the folders are scanned.
        self.unusable_dirs: list[Path] = []

        if not self.outputs_root.exists():
            raise FileNotFoundError(f"Pasta não encontrada: {self.outputs_root}")

    # =====================================================
    # PATHS
    # =====================================================
    # Marker written by the processor for a document whose text could not be
    # used — an image-only scan, or a text layer whose font has no usable
    # character map (see scanned_pdf_detector).
    UNUSABLE_MARKER = "NAO_UTILIZAVEL.json"
    LEGACY_UNUSABLE_MARKER = "DIGITALIZACAO.json"

    def _find_doc_dirs(self) -> list[Path]:
        """Document folders eligible for the analysis.

        Unusable documents are held apart rather than counted. They produce no
        entities by construction, so leaving them in `total_documents_processed`
        would mean every percentage is computed over a denominator that
        includes documents the pipeline never had a chance to read — a
        systematic understatement of coverage.

        They are reported separately in the summary, so the exclusion is
        visible rather than silent.
        """
        todas = sorted([p for p in self.outputs_root.iterdir() if p.is_dir()])

        self.unusable_dirs = [
            p for p in todas
            if (p / self.UNUSABLE_MARKER).exists()
            or (p / self.LEGACY_UNUSABLE_MARKER).exists()
        ]
        excluidas = {p.name for p in self.unusable_dirs}
        return [p for p in todas if p.name not in excluidas]

    def _unusable_rows(self) -> list[dict]:
        """Um registo por documento excluído, com a causa.

        Excluir dos denominadores não pode significar fazer desaparecer. Um
        RCM que o pipeline não consegue ler é, quase sempre, um RCM que um
        profissional também não lê bem: digitalização sem texto, fonte sem
        mapa de caracteres, títulos sem numeração nem distinção tipográfica.
        A lista destes casos é um resultado sobre a qualidade editorial dos
        rótulos, não apenas uma limitação da ferramenta.
        """
        CAUSAS = {
            "scanned": "Digitalização sem camada de texto",
            "corrupt": "Fonte sem mapa de caracteres — texto ilegível",
            "sem_estrutura": "Títulos sem numeração nem distinção tipográfica",
        }
        linhas = []
        for pasta in self.unusable_dirs:
            dados = {}
            for nome in (self.UNUSABLE_MARKER, self.LEGACY_UNUSABLE_MARKER):
                caminho = pasta / nome
                if caminho.exists():
                    try:
                        dados = json.loads(caminho.read_text(encoding="utf-8"))
                    except Exception:  # noqa: BLE001
                        dados = {}
                    break
            estado = str(dados.get("status") or "desconhecido")
            linhas.append({
                "documento": pasta.name,
                "estado": estado,
                "causa": CAUSAS.get(estado, "Não classificado"),
                "detalhe": str(dados.get("reason") or ""),
                "caracteres": dados.get("chars", ""),
            })
        return sorted(linhas, key=lambda x: (x["estado"], x["documento"]))

    def _pgx_dir(self, doc_dir: Path) -> Path:
        return doc_dir / self.PGX_OUTPUTS_DIRNAME

    def _unique_path_for_doc(self, doc_dir: Path) -> Path | None:
        candidates = [
            self._pgx_dir(doc_dir) / self.UNIQUE_FILENAME,
            doc_dir / self.UNIQUE_FILENAME,
        ]

        for path in candidates:
            if path.exists():
                return path

        return None

    def _extended_path_for_doc(self, doc_dir: Path) -> Path | None:
        candidates = []

        for filename in self.EXTENDED_FILENAMES:
            candidates.append(self._pgx_dir(doc_dir) / filename)
            candidates.append(doc_dir / filename)

        for path in candidates:
            if path.exists():
                return path

        return None

    def _iter_section_jsons(self, doc_dir: Path) -> list[Path]:
        out = []

        for path in sorted(doc_dir.glob("*.json")):
            if path.name in self.SPECIAL_FILES:
                continue

            data = self._read_json_safe(path)
            avaliacao = self._safe_dict(data.get("avaliacao_farmacogenomica"))

            if avaliacao:
                out.append(path)

        return out

    # =====================================================
    # IO / HELPERS
    # =====================================================
    def _read_json_safe(self, path: Path | None) -> dict:
        if path is None:
            return {}

        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return {}

    def _safe_list(self, value: Any) -> list:
        return value if isinstance(value, list) else []

    def _safe_dict(self, value: Any) -> dict:
        return value if isinstance(value, dict) else {}

    def _pct(self, part: int | float, total: int | float) -> float:
        if not total:
            return 0.0
        return round((part / total) * 100, 2)

    def _bool_value(self, value: Any) -> bool:
        if isinstance(value, bool):
            return value

        if isinstance(value, int):
            return value > 0

        if isinstance(value, str):
            return value.strip().lower() in {"sim", "true", "yes", "1"}

        return False

    def _to_int(self, value: Any, default: int = 0) -> int:
        try:
            return int(value)
        except Exception:
            return default

    def _join_values(self, values: list[str] | set[str]) -> str:
        clean = [str(v).strip() for v in values if str(v).strip()]
        return " ; ".join(sorted(set(clean)))

    def _format_counter(self, counter: Counter) -> str:
        if not counter:
            return ""
        return " ; ".join(f"{k}={v}" for k, v in sorted(counter.items()))

    def _entity_total_mentions(self, item: dict, default: int = 1) -> int:
        for key in ["total_mentions", "mentions", "count", "occurrences"]:
            value = item.get(key)
            try:
                n = int(value)
                if n > 0:
                    return n
            except Exception:
                continue

        return default

    # HLA na forma anterior à revisão da nomenclatura: HLA-B*5701.
    #
    # A forma actual separa os campos com dois pontos, HLA-B*57:01, e os RCM
    # portugueses usam as duas. Sem fundir, o mesmo alelo conta como duas
    # entidades distintas — o HLA-B*57:01 do abacavir aparecia repartido
    # entre 11 documentos numa grafia e 10 na outra.
    #
    # A fusão é feita aqui, no funil por onde todas as entidades passam, e
    # não sobre as contagens já agregadas: as menções somar-se-iam bem, mas
    # os documentos não, porque um documento que use as duas grafias seria
    # contado duas vezes.
    #
    # Só o caso inequívoco de quatro dígitos é convertido, que é 2+2 em
    # todas as designações correntes. Cinco ou mais dígitos admitem mais do
    # que uma leitura e ficam como estão, para não inventar.
    _HLA_SEM_DOIS_PONTOS = re.compile(r"^(HLA-[A-Z]+\d*)\*(\d{4})$", re.IGNORECASE)

    @classmethod
    def _canonizar_entidade(cls, valor: str) -> str:
        m = cls._HLA_SEM_DOIS_PONTOS.match(valor)
        if m:
            locus, digitos = m.group(1).upper(), m.group(2)
            return f"{locus}*{digitos[:2]}:{digitos[2:]}"
        return valor

    # Gene a que uma entidade pertence: CYP2C19*2 e CYP2C19*1*2 pertencem a
    # CYP2C19, HLA-B*15:02 a HLA-B, SLCO1B1c.521CC a SLCO1B1.
    #
    # Uma entidade sem gene identificável — as bases soltas "CC", "TT" —
    # devolve vazio e não é atribuída a ninguém.
    _GENE_HGVS = re.compile(r"^([A-Z][A-Z0-9-]{2,})c\.\d+", re.IGNORECASE)
    _GENE_BASES_FIM = re.compile(r"^([A-Z][A-Z0-9-]{2,}?)[ACGT]{2}$", re.IGNORECASE)

    @classmethod
    def _gene_de_entidade(cls, entidade: str) -> str:
        e = (entidade or "").strip()
        if not e:
            return ""
        if "*" in e:
            return e.split("*", 1)[0].upper()
        m = cls._GENE_HGVS.match(e)
        if m:
            return m.group(1).upper()
        m = cls._GENE_BASES_FIM.match(e)
        if m:
            return m.group(1).upper()
        return ""

    def _entity_value(self, item: dict) -> str:
        valor = str(
            item.get("entity")
            or item.get("variant")
            or item.get("symbol")
            or item.get("rsid")
            or ""
        ).strip()
        return self._canonizar_entidade(valor)

    def _gene_value(self, item: dict) -> str:
        return str(
            item.get("symbol")
            or item.get("entity")
            or item.get("gene")
            or ""
        ).strip()

    # =====================================================
    # DOCUMENT_UNIQUE SCHEMA
    # =====================================================
    def _document_block(self, unique_data: dict) -> dict:
        doc = self._safe_dict(unique_data.get("documento"))

        return {
            "nome_medicamento": doc.get("nome_medicamento", ""),
            "substancia_ativa": doc.get("substancia_ativa", ""),
            "aliases_substancia_ativa": self._safe_list(doc.get("aliases_substancia_ativa")),
            "grupo_farmacoterapeutico": doc.get("grupo_farmacoterapeutico", ""),
            "grupo_farmacoterapeutico_codigo": doc.get("grupo_farmacoterapeutico_codigo", ""),
            "codigo_atc": doc.get("codigo_atc", ""),
        }

    def _extraction_quality(self, unique_data: dict) -> tuple[int, int]:
        """Dois indicadores de saúde da extração deste documento.

        Returns:
            (falhas, nao_verbatim)

            falhas       - blocos cuja resposta do modelo não era JSON
                           utilizável. O conteúdo PGx desses blocos perdeu-se,
                           logo as contagens podem estar SUBESTIMADAS.
            nao_verbatim - excertos devolvidos pelo modelo que não constam
                           literalmente do bloco de origem, ou seja
                           parafraseados. Ficam fora da contagem de entidades.

        Ambos a 0 significa extração limpa. Valores altos indicam documentos
        que merecem revisão manual antes de os números serem citados.
        """
        block = self._safe_dict(unique_data.get("qualidade_extracao"))
        return (
            self._to_int(block.get("falhas_extracao")),
            self._to_int(block.get("excertos_recompostos"))
            + self._to_int(block.get("excertos_nao_encontrados")),
        )

    def _origin_block(self, unique_data: dict) -> dict:
        """Repartição das menções entre texto corrido e tabelas.

        As tabelas dos RCMs eram removidas do `content` pelo MarkdownProcessor
        e nunca lidas pelo avaliador. No subset de 17 documentos isso custava
        os seis diplótipos CYP2C9 do siponimod, com frequências populacionais e
        impacto na exposição.

        Nota sobre os totais: as entidades únicas por origem podem somar mais
        do que o total do documento, porque a mesma entidade pode aparecer no
        texto e numa tabela. As menções, essas, somam certo.
        """
        block = self._safe_dict(unique_data.get("origem_entidades"))
        texto = self._safe_dict(block.get("texto"))
        tabelas = self._safe_dict(block.get("tabelas"))
        return {
            "texto_unicas": self._to_int(texto.get("entidades_unicas")),
            "texto_mencoes": self._to_int(texto.get("mencoes")),
            "tabelas_unicas": self._to_int(tabelas.get("entidades_unicas")),
            "tabelas_mencoes": self._to_int(tabelas.get("mencoes")),
            "exclusivas_de_tabelas": self._safe_list(block.get("exclusivas_de_tabelas")),
        }

    def _pipeline_block(self, unique_data: dict) -> dict:
        """Proveniência: versão do pipeline e configuração do modelo."""
        return self._safe_dict(unique_data.get("pipeline"))

    def _identifiers_block(self, unique_data: dict) -> dict:
        ids = self._safe_dict(unique_data.get("identificadores_substancia"))

        return {
            "mesh_id": ids.get("mesh_id", ""),
            "drugbank_id": ids.get("drugbank_id", ""),
            "chebi_id": ids.get("chebi_id", ""),
            "umls_cui": ids.get("umls_cui") or ids.get("mrconso_cui", ""),
        }

    def _counts_block(self, unique_data: dict, extended_data: dict | None = None) -> dict:
        extended_data = extended_data or {}

        return (
            self._safe_dict(unique_data.get("contagens"))
            or self._safe_dict(extended_data.get("contagens_documento"))
            or self._safe_dict(unique_data.get("contagens_finais"))
        )

    def _entities(self, unique_data: dict, category: str) -> list[dict]:
        entidades = self._safe_dict(unique_data.get("entidades"))

        value = entidades.get(category)
        if isinstance(value, list):
            return value

        entidades_unicas = self._safe_dict(unique_data.get("entidades_unicas"))
        value = entidades_unicas.get(category)
        if isinstance(value, list):
            return value

        legacy_map = {
            "genes": "document_unique_genes",
            "star_alleles": "document_unique_star_alleles",
            "diplotypes": "document_unique_diplotypes",
            "rsids": "document_unique_rsids",
        }

        value = unique_data.get(legacy_map.get(category, ""))
        return value if isinstance(value, list) else []

    def _unique_entities(self, unique_data: dict) -> dict[str, list[str]]:
        result = {
            "genes": [],
            "star_alleles": [],
            "diplotypes": [],
            "rsids": [],
        }

        for item in self._entities(unique_data, "genes"):
            if isinstance(item, dict):
                value = self._gene_value(item)
                if value:
                    result["genes"].append(value)

        for category in ["star_alleles", "diplotypes", "rsids"]:
            for item in self._entities(unique_data, category):
                if isinstance(item, dict):
                    value = self._entity_value(item)
                    if value:
                        result[category].append(value)

        return {
            key: sorted(set(values))
            for key, values in result.items()
        }

    def _has_entities(self, unique_data: dict, extended_data: dict | None = None) -> bool:
        counts = self._counts_block(unique_data, extended_data)

        if self._to_int(counts.get("entidades_unicas_total")) > 0:
            return True

        if self._to_int(counts.get("entidades_mencoes_total")) > 0:
            return True

        entities = self._unique_entities(unique_data)

        return any([
            len(entities["genes"]) > 0,
            len(entities["star_alleles"]) > 0,
            len(entities["diplotypes"]) > 0,
            len(entities["rsids"]) > 0,
        ])

    # =====================================================
    # GUIDELINES / CLINPGX
    # =====================================================
    def _entity_guideline_ids(self, item: dict) -> list[str]:
        ids = self._safe_dict(item.get("ids"))

        values = ids.get("guideline_ids")
        if isinstance(values, list):
            return [str(v).strip() for v in values if str(v).strip()]

        values = item.get("guideline_ids")
        if isinstance(values, list):
            return [str(v).strip() for v in values if str(v).strip()]

        guideline_matches = item.get("guideline_matches")
        if isinstance(guideline_matches, list):
            out = []
            for match in guideline_matches:
                if isinstance(match, dict):
                    guideline_id = str(match.get("guideline_id") or "").strip()
                    if guideline_id:
                        out.append(guideline_id)
            return out

        return []

    def _entity_has_guideline(self, item: dict) -> bool:
        """A entidade está associada a uma guideline annotation do ClinPGx?

        Só conta evidência proveniente de `guidelineAnnotations`. Não conta:

          * `genes_tsv` / `is_known_pgx_gene` — presença em genes.tsv significa
            que o gene é farmacogeneticamente relevante em abstrato, não que
            exista uma guideline para este par fármaco-gene;
          * `clinicalVariants_tsv` / `has_clinical_variant_evidence` — uma
            anotação de variante clínica é evidência, mas não é uma guideline.

        Manter estes três níveis separados é o que permite reportar
        "gene conhecido", "com evidência de variante" e "com guideline" como
        colunas distintas em vez de os fundir num único sim/não.
        """
        if self._entity_guideline_ids(item):
            return True

        evidence_flags = self._safe_dict(item.get("evidence_flags"))
        if evidence_flags.get("has_guideline"):
            return True

        source_matches = self._safe_dict(item.get("source_matches"))
        if source_matches.get("clinpgx_guideline_json"):
            return True

        # Removido: `item.get("matches")`. Era um teste genérico por uma chave
        # que nenhum output usa hoje, mas que — se algum dia contivesse os
        # clinical_variant_matches — classificaria evidência de variante como
        # guideline. A distinção é precisamente o que aqui interessa preservar.
        return False

    # =====================================================
    # COBERTURA DE GUIDELINES
    # =====================================================
    #
    # Uma guideline do ClinPGx é um par FÁRMACO-GENE. Não é uma anotação em
    # genes.tsv (isso diz que o gene é farmacogeneticamente relevante em
    # abstrato) nem em clinicalVariants.tsv (isso é evidência de variante).
    # Só `guidelineAnnotations` conta aqui.
    #
    # A pergunta que estas funções respondem:
    #   dos documentos cuja substância ativa tem pelo menos uma guideline,
    #   quais é que mencionam de facto o(s) gene(s) dessa guideline?
    #
    # Três níveis por par documento×guideline:
    #   coberta   -> o RCM menciona TODOS os genes da guideline
    #   parcial   -> menciona alguns (só possível em guidelines multi-gene)
    #   ausente   -> não menciona nenhum
    #
    # No snapshot ClinPGx em uso, 190 das 217 guidelines têm um único gene,
    # pelo que "parcial" aplica-se apenas às 27 multi-gene.

    def _observed_gene_symbols(self, unique_data: dict) -> set[str]:
        """Genes que o documento referencia, direta ou indiretamente.

        Um RCM que escreve "CYP2D6*4" está a referenciar CYP2D6, mesmo que
        nunca mencione o símbolo isolado. Por isso somam-se aos genes extraídos
        como entidade os genes atribuídos a star alleles, diplótipos e rsIDs.
        Contar apenas os genes isolados subestimaria a cobertura precisamente
        nos documentos mais informativos, que são os que dão o alelo concreto.
        """
        observed: set[str] = set()

        for item in self._entities(unique_data, "genes"):
            if isinstance(item, dict):
                value = self._gene_value(item)
                if value:
                    observed.add(self._norm_symbol(value))

        for category in ("star_alleles", "diplotypes", "rsids"):
            for item in self._entities(unique_data, category):
                if not isinstance(item, dict):
                    continue
                # star alleles e rsIDs trazem "gene"; diplótipos trazem "genes"
                gene = item.get("gene")
                if gene:
                    observed.add(self._norm_symbol(str(gene)))
                for gene in self._safe_list(item.get("genes")):
                    if gene:
                        observed.add(self._norm_symbol(str(gene)))
                # último recurso: prefixo antes do '*' (ex. "CYP2C9*3")
                entity = self._entity_value(item)
                if "*" in entity:
                    prefix = entity.split("*", 1)[0].strip()
                    if prefix:
                        observed.add(self._norm_symbol(prefix))

        observed.discard("")
        return observed

    @staticmethod
    def _norm_symbol(value: str) -> str:
        return re.sub(r"\s+", "", str(value or "")).upper()

    def _substance_guideline_matches(self, extended_data: dict) -> list[dict]:
        """Guidelines cujo relatedChemicals correspondeu à substância ativa."""
        refs = self._safe_dict(extended_data.get("active_substance_external_references"))
        clinpgx = self._safe_dict(refs.get("clinpgx"))
        return [m for m in self._safe_list(clinpgx.get("matches")) if isinstance(m, dict)]

    def _guideline_expected_genes(self, match: dict) -> list[str]:
        """Símbolos dos genes que a guideline exige."""
        out = []
        for gene in self._safe_list(match.get("related_genes")):
            if isinstance(gene, dict):
                symbol = self._norm_symbol(gene.get("symbol") or "")
                if symbol:
                    out.append(symbol)
        return sorted(set(out))

    def _guideline_coverage_rows(
        self,
        doc_name: str,
        active_substance: str,
        unique_data: dict,
        extended_data: dict,
    ) -> list[dict]:
        """Uma linha por par documento×guideline da substância ativa."""
        observed = self._observed_gene_symbols(unique_data)
        rows = []

        for match in self._substance_guideline_matches(extended_data):
            expected = self._guideline_expected_genes(match)
            if not expected:
                # Guideline sem genes associados: não é um par fármaco-gene,
                # logo não é avaliável em termos de cobertura.
                continue

            covered = sorted(set(expected) & observed)
            missing = sorted(set(expected) - observed)

            if not covered:
                status = "ausente"
            elif missing:
                status = "parcial"
            else:
                status = "coberta"

            rows.append({
                "document": doc_name,
                "active_substance": active_substance,
                "guideline_id": match.get("guideline_id") or "",
                "guideline_name": match.get("guideline_name") or "",
                "source": match.get("source") or "",
                "expected_genes": expected,
                "covered_genes": covered,
                "missing_genes": missing,
                "n_expected": len(expected),
                "n_covered": len(covered),
                "status": status,
            })

        return rows

    def _substance_has_guideline(self, extended_data: dict) -> bool:
        """Existe alguma guideline ClinPGx para a substância ativa deste RCM?

        Depende apenas do nome da substância ter correspondido a um registo do
        ClinPGx. NÃO implica que o RCM mencione o gene relevante — é uma
        propriedade do fármaco, não do documento.

        Útil como denominador ("de quantos fármacos com guideline é que o RCM
        chega a mencionar o gene?"), mas não serve como resposta a
        "este documento tem informação de guideline".
        """
        refs = self._safe_dict(extended_data.get("active_substance_external_references"))
        clinpgx = self._safe_dict(refs.get("clinpgx"))
        return bool(clinpgx.get("found"))

    def _has_guideline(self, unique_data: dict, extended_data: dict) -> bool:
        """O documento contém um par fármaco-gene coberto por uma guideline?

        Definição farmacogenomicamente correta, alinhada com CPIC/DPWG: uma
        guideline só se aplica quando a substância E o gene relevante estão
        ambos presentes. Uma entidade (gene, alelo, diplótipo ou rsID) detetada
        no RCM tem de constar de uma guideline annotation cujos
        `relatedChemicals` já correspondiam à substância ativa.

        Antes bastava `active_substance_external_references.clinpgx.found` —
        ou seja, o nome da substância ter correspondido a um registo ClinPGx,
        independentemente do RCM mencionar seja o que for. No corpus de 643
        RCMs isso produzia 49 documentos "com guideline" e ZERO conteúdo PGx,
        36% dos 135 contabilizados.

        Nota: presença em genes.tsv ou em clinicalVariants.tsv não conta aqui.
        São níveis de evidência distintos e reportados em colunas próprias.
        """
        # Via primária: existe alguma guideline da substância cujo gene o
        # documento menciona. É a definição literal do par fármaco-gene e não
        # depende do output compacto trazer guideline_ids — foi precisamente a
        # omissão desse campo nos genes que escondeu o problema durante tanto
        # tempo.
        rows = self._guideline_coverage_rows("", "", unique_data, extended_data)
        if any(row["status"] != "ausente" for row in rows):
            return True

        # Via secundária: as marcas de guideline já resolvidas pelo avaliador.
        # Mantida para documentos cujo Extended_PGx_Analysis não esteja
        # disponível.
        counts = self._counts_block(unique_data, extended_data)

        for key in [
            "genes_com_guideline_total",
            "star_alleles_com_guideline_total",
            "diplotipos_com_guideline_total",
            "rsids_com_guideline_total",
        ]:
            if self._to_int(counts.get(key)) > 0:
                return True

        for category in ["genes", "star_alleles", "diplotypes", "rsids"]:
            for item in self._entities(unique_data, category):
                if isinstance(item, dict) and self._entity_has_guideline(item):
                    return True

        return False

    def _guideline_ids_for_doc(self, unique_data: dict, extended_data: dict) -> list[str]:
        """IDs das guidelines efetivamente suportadas por entidades do documento.

        Antes juntava também todos os IDs vindos do match ao nível da
        substância, pelo que a coluna "Guideline IDs" vinha preenchida mesmo em
        documentos sem uma única entidade PGx. Esses continuam disponíveis em
        `_substance_guideline_ids`, numa coluna própria.
        """
        out = []

        for category in ["genes", "star_alleles", "diplotypes", "rsids"]:
            for item in self._entities(unique_data, category):
                if isinstance(item, dict):
                    out.extend(self._entity_guideline_ids(item))

        return sorted(set(out))

    def _substance_guideline_ids(self, extended_data: dict) -> list[str]:
        """Guidelines que existem para a substância, mencione-as o RCM ou não."""
        refs = self._safe_dict(extended_data.get("active_substance_external_references"))
        clinpgx = self._safe_dict(refs.get("clinpgx"))

        out = []
        for match in self._safe_list(clinpgx.get("matches")):
            if isinstance(match, dict):
                guideline_id = str(match.get("guideline_id") or "").strip()
                if guideline_id:
                    out.append(guideline_id)

        return sorted(set(out))

    def _guideline_sources_for_doc(self, extended_data: dict) -> str:
        refs = self._safe_dict(extended_data.get("active_substance_external_references"))
        clinpgx = self._safe_dict(refs.get("clinpgx"))

        sources = []

        for match in self._safe_list(clinpgx.get("matches")):
            if isinstance(match, dict):
                source = str(match.get("source") or "").strip()
                if source:
                    sources.append(source)

        return self._join_values(sources)

    # =====================================================
    # SECÇÕES
    # =====================================================
    def _section_has_pgx(self, section_data: dict) -> bool:
        avaliacao = self._safe_dict(section_data.get("avaliacao_farmacogenomica"))

        value = (
            avaliacao.get("contem_farmacogenomica")
            or avaliacao.get("has_pgx")
            or section_data.get("has_pgx")
        )

        return self._bool_value(value)

    def _has_pgx(self, unique_data: dict, extended_data: dict, section_paths: list[Path]) -> bool:
        classificacao = self._safe_dict(extended_data.get("classificacao"))

        if "has_pgx" in classificacao:
            return self._bool_value(classificacao.get("has_pgx"))

        if "has_pgx" in extended_data:
            return self._bool_value(extended_data.get("has_pgx"))

        for path in section_paths:
            data = self._read_json_safe(path)
            if self._section_has_pgx(data):
                return True

        if self._has_entities(unique_data, extended_data):
            return True

        return False

    # Título canónico por número de secção.
    #
    # A chave de agrupamento era "número + título tal como vem do RCM", e o
    # título varia entre documentos: ortografia antiga ("Interacções ... de
    # interacção"), hífen ("Contra-indicações"), maiúsculas ("PROPRIEDADES
    # FARMACOCINÉTICAS") e, quando a conversão duplica a linha, títulos
    # repetidos dentro de si mesmos. A folha ficava com a 4.5 partida em três
    # linhas e a 4.3 em duas, e ninguém somava certo.
    #
    # O número é o identificador estável — é normativo e igual em todos os
    # RCMs. O título passa a ser decorativo e vem desta tabela.
    TITULOS_CANONICOS = {
        "4.1": "Indicações terapêuticas",
        "4.2": "Posologia e modo de administração",
        "4.3": "Contraindicações",
        "4.4": "Advertências e precauções especiais de utilização",
        "4.5": "Interações medicamentosas e outras formas de interação",
        "4.6": "Fertilidade, gravidez e aleitamento",
        "4.7": "Efeitos sobre a capacidade de conduzir e utilizar máquinas",
        "4.8": "Efeitos indesejáveis",
        "4.9": "Sobredosagem",
        "5.1": "Propriedades farmacodinâmicas",
        "5.2": "Propriedades farmacocinéticas",
        "5.3": "Dados de segurança pré-clínica",
    }

    def _section_key(self, section_data: dict, fallback_name: str) -> str:
        """Chave de agrupamento da secção: o número, com título canónico.

        Sem isto, "4.3 Contraindicações" e "4.3 Contra-indicações" contam como
        duas secções diferentes.
        """
        number = str(
            section_data.get("numero")
            or section_data.get("section_number")
            or section_data.get("number")
            or ""
        ).strip()

        if number in self.TITULOS_CANONICOS:
            return f"{number} {self.TITULOS_CANONICOS[number]}"

        # Secção fora do padrão normativo: mantém-se o que o documento diz,
        # para não esconder um caso que precise de ser visto.
        title = str(
            section_data.get("titulo")
            or section_data.get("section_title")
            or section_data.get("title")
            or fallback_name
        ).strip()
        return f"{number} {title}".strip()

    def _extract_section_entities(self, section_data: dict) -> dict[str, Counter]:
        result = {
            "genes": Counter(),
            "star_alleles": Counter(),
            "diplotypes": Counter(),
            "rsids": Counter(),
        }

        avaliacao = self._safe_dict(section_data.get("avaliacao_farmacogenomica"))
        entidades = self._safe_dict(avaliacao.get("pgx_entidades"))

        for item in self._safe_list(entidades.get("genes")):
            if isinstance(item, dict):
                value = self._gene_value(item)
                if value:
                    result["genes"][value] += self._entity_total_mentions(item)

        for item in self._safe_list(entidades.get("star_alleles")):
            if isinstance(item, dict):
                value = self._entity_value(item)
                if value:
                    result["star_alleles"][value] += self._entity_total_mentions(item)

        for item in self._safe_list(entidades.get("diplotypes")):
            if isinstance(item, dict):
                value = self._entity_value(item)
                if value:
                    result["diplotypes"][value] += self._entity_total_mentions(item)

        for item in self._safe_list(entidades.get("rsids")):
            if isinstance(item, dict):
                value = self._entity_value(item)
                if value:
                    result["rsids"][value] += self._entity_total_mentions(item)

        return result

    def _section_normalization_rows(
        self,
        section_data: dict,
        doc_name: str,
        active_substance: str,
        section_label: str,
    ) -> list[dict]:
        rows = []

        avaliacao = self._safe_dict(section_data.get("avaliacao_farmacogenomica"))
        normalizacao = self._safe_dict(avaliacao.get("normalizacao"))

        for key, label in [
            ("nome_extenso_para_simbolo", "nome_extenso_para_simbolo"),
            ("normalizacao_tecnica", "normalizacao_tecnica"),
            ("star_alleles_incompletos", "star_alleles_incompletos"),
        ]:
            for note in self._safe_list(normalizacao.get(key)):
                if not isinstance(note, dict):
                    continue

                rows.append({
                    "document": doc_name,
                    "active_substance": active_substance,
                    "section": section_label,
                    "tipo": label,
                    "original_text": note.get("original_text", ""),
                    "normalized_text": (
                        note.get("normalized_symbol")
                        or note.get("normalized_text")
                        or note.get("normalized_variant")
                        or ""
                    ),
                    "entity_type": note.get("entity_type", ""),
                    "total_mentions": note.get("total_mentions", 1),
                    "source": "section_json",
                })

        return rows

    # =====================================================
    # NORMALIZAÇÕES
    # =====================================================
    def _extended_normalization_rows(
        self,
        extended_data: dict,
        doc_name: str,
        active_substance: str,
    ) -> list[dict]:
        rows = []
        normalizacao = self._safe_dict(extended_data.get("normalizacao"))

        for key, label in [
            ("nome_extenso_para_simbolo", "nome_extenso_para_simbolo"),
            ("normalizacao_tecnica", "normalizacao_tecnica"),
            ("star_alleles_incompletos", "star_alleles_incompletos"),
        ]:
            for note in self._safe_list(normalizacao.get(key)):
                if not isinstance(note, dict):
                    continue

                rows.append({
                    "document": doc_name,
                    "active_substance": active_substance,
                    "section": note.get("section", ""),
                    "tipo": label,
                    "original_text": note.get("original_text", ""),
                    "normalized_text": (
                        note.get("normalized_symbol")
                        or note.get("normalized_text")
                        or note.get("normalized_variant")
                        or ""
                    ),
                    "entity_type": note.get("entity_type", ""),
                    "total_mentions": note.get("total_mentions", 1),
                    "source": "Extended_PGx_Analysis",
                })

        return rows

    def _aggregate_normalization_rows_if_needed(
        self,
        unique_data: dict,
        extended_data: dict,
        doc_name: str,
        active_substance: str,
        existing_rows: list[dict],
    ) -> list[dict]:
        rows = list(existing_rows)
        counts = self._counts_block(unique_data, extended_data)

        expected = {
            "nome_extenso_para_simbolo": self._to_int(
                counts.get("genes_normalizados_nome_extenso_para_simbolo_total")
            ),
            "normalizacao_tecnica": self._to_int(
                counts.get("normalizacoes_tecnicas_total")
            ),
            "star_alleles_incompletos": self._to_int(
                counts.get("star_alleles_incompletos_total")
            ),
        }

        for tipo, total in expected.items():
            if total <= 0:
                continue

            already = sum(
                self._to_int(row.get("total_mentions"), 1)
                for row in rows
                if row.get("tipo") == tipo
            )

            if already == 0:
                rows.append({
                    "document": doc_name,
                    "active_substance": active_substance,
                    "section": "",
                    "tipo": tipo,
                    "original_text": "(detalhe não guardado no JSON)",
                    "normalized_text": "(detalhe não guardado no JSON)",
                    "entity_type": "gene" if tipo == "nome_extenso_para_simbolo" else "",
                    "total_mentions": total,
                    "source": "contagem_agregada",
                })

        return rows

    # =====================================================
    # ANÁLISE
    # =====================================================
    def analyse(self) -> dict:
        doc_dirs = self._find_doc_dirs()

        documents = []
        farmaco_entidades = []
        normalizacao_rows = []
        entidades_clinpgx_rows = []
        section_rows = []

        total_docs = 0
        docs_with_unique = 0
        docs_with_pgx = 0
        docs_without_pgx = 0
        docs_with_pgx_entities = 0
        docs_with_pgx_no_entities = 0
        docs_with_guideline = 0
        docs_without_guideline = 0
        docs_with_guideline_and_pgx = 0
        docs_with_guideline_without_pgx = 0
        # Nível da substância: existe guideline para o fármaco, mencione-a o
        # RCM ou não. O contraste com docs_with_guideline mede a lacuna
        # regulamentar — fármacos com PGx conhecida cujo RCM a omite.
        docs_substance_with_guideline = 0
        docs_substance_guideline_not_in_label = 0

        # Cobertura: uma linha por par documento×guideline
        guideline_coverage_rows: list[dict] = []
        coverage_status_counts = Counter()
        # Pares (documento, gene) sem repetição — ver comentário no ciclo.
        distinct_expected: set[tuple[str, str]] = set()
        distinct_covered: set[tuple[str, str]] = set()
        genes_expected_total = 0
        genes_covered_total = 0
        docs_fully_covered = 0
        docs_partially_covered = 0
        docs_not_covered = 0
        docs_with_guideline_missing_genes = 0
        origin_rows: list[dict] = []
        total_origin = Counter()

        total_extraction_failures = 0
        docs_with_extraction_failures = 0
        total_non_verbatim = 0
        docs_with_non_verbatim = 0
        pipeline_versions = Counter()
        llm_configs = Counter()
        # gene -> (nº pares em que era esperado, nº em que foi mencionado)
        # Por gene: posições de guideline (com repetição) e documentos
        # distintos. Ver comentário no ciclo — não são a mesma coisa.
        gene_expected_positions = Counter()
        gene_covered_positions = Counter()
        gene_expected_docs_set: defaultdict[str, set] = defaultdict(set)
        gene_covered_docs_set: defaultdict[str, set] = defaultdict(set)

        # Gene agregado: menções do símbolo mais as dos seus alelos e
        # diplótipos. Ver _gene_de_entidade. Os documentos são acumulados
        # como conjunto e não como soma, porque um rótulo que cita
        # CYP2C19*2 cita quase sempre CYP2C19 também.
        gene_mais_docs: dict[str, set[str]] = defaultdict(set)
        gene_mais_mencoes = Counter()

        global_unique_gene_docs = Counter()
        global_unique_allele_docs = Counter()
        global_unique_diplotype_docs = Counter()
        global_unique_rsid_docs = Counter()

        global_gene_mentions = Counter()
        global_allele_mentions = Counter()
        global_diplotype_mentions = Counter()
        global_rsid_mentions = Counter()

        genes_without_tsv_docs = defaultdict(set)

        section_stats = defaultdict(lambda: {
            "documents": set(),
            "documents_with_pgx": set(),
            "gene_mentions": 0,
            "star_allele_mentions": 0,
            "diplotype_mentions": 0,
            "rsid_mentions": 0,
            "total_mentions": 0,
        })

        group_stats = defaultdict(lambda: {
            "documents": set(),
            "documents_with_pgx": set(),
            "documents_with_pgx_entities": set(),
            "documents_with_guideline": set(),
            "gene_mentions": 0,
            "unique_genes": set(),
        })

        # código do grupo -> rótulo textual mais completo observado
        group_labels: dict[str, str] = {}

        total_model_normalizations = 0
        total_technical_normalizations = 0
        total_incomplete_star_alleles = 0
        total_gene_mentions_with_guideline = 0
        total_entity_mentions_with_guideline = 0

        for doc_dir in doc_dirs:
            total_docs += 1
            doc_name = doc_dir.name

            unique_path = self._unique_path_for_doc(doc_dir)
            extended_path = self._extended_path_for_doc(doc_dir)
            section_paths = self._iter_section_jsons(doc_dir)

            if unique_path is None:
                documents.append({
                    "document": doc_name,
                    "status": "missing_document_unique",
                })
                continue

            docs_with_unique += 1

            unique_data = self._read_json_safe(unique_path)
            extended_data = self._read_json_safe(extended_path)

            doc = self._document_block(unique_data)
            ids = self._identifiers_block(unique_data)
            counts = self._counts_block(unique_data, extended_data)
            unique_entities = self._unique_entities(unique_data)

            medicamento = doc["nome_medicamento"]
            active_substance = doc["substancia_ativa"]
            aliases = doc["aliases_substancia_ativa"]
            grupo = doc["grupo_farmacoterapeutico"]
            grupo_codigo = doc["grupo_farmacoterapeutico_codigo"]
            atc = doc["codigo_atc"]

            doc_mentions = {
                "genes": Counter(),
                "star_alleles": Counter(),
                "diplotypes": Counter(),
                "rsids": Counter(),
            }

            doc_normalization_rows = []

            for section_path in section_paths:
                section_data = self._read_json_safe(section_path)
                section = self._section_key(section_data, section_path.stem)
                section_entities = self._extract_section_entities(section_data)

                gene_mentions = sum(section_entities["genes"].values())
                star_mentions = sum(section_entities["star_alleles"].values())
                diplotype_mentions = sum(section_entities["diplotypes"].values())
                rsid_mentions = sum(section_entities["rsids"].values())

                total_section_mentions = (
                    gene_mentions
                    + star_mentions
                    + diplotype_mentions
                    + rsid_mentions
                )

                if total_section_mentions > 0:
                    doc_mentions["genes"].update(section_entities["genes"])
                    doc_mentions["star_alleles"].update(section_entities["star_alleles"])
                    doc_mentions["diplotypes"].update(section_entities["diplotypes"])
                    doc_mentions["rsids"].update(section_entities["rsids"])

                    section_stats[section]["documents"].add(doc_name)
                    section_stats[section]["documents_with_pgx"].add(doc_name)
                    section_stats[section]["gene_mentions"] += gene_mentions
                    section_stats[section]["star_allele_mentions"] += star_mentions
                    section_stats[section]["diplotype_mentions"] += diplotype_mentions
                    section_stats[section]["rsid_mentions"] += rsid_mentions
                    section_stats[section]["total_mentions"] += total_section_mentions

                    section_rows.append({
                        "document": doc_name,
                        "section": section,
                        "gene_mentions": gene_mentions,
                        "star_allele_mentions": star_mentions,
                        "diplotype_mentions": diplotype_mentions,
                        "rsid_mentions": rsid_mentions,
                        "total_mentions": total_section_mentions,
                        "genes": self._format_counter(section_entities["genes"]),
                        "star_alleles": self._format_counter(section_entities["star_alleles"]),
                        "diplotypes": self._format_counter(section_entities["diplotypes"]),
                        "rsids": self._format_counter(section_entities["rsids"]),
                    })

                doc_normalization_rows.extend(
                    self._section_normalization_rows(
                        section_data=section_data,
                        doc_name=doc_name,
                        active_substance=active_substance,
                        section_label=section,
                    )
                )

            has_entities = self._has_entities(unique_data, extended_data)
            has_pgx = self._has_pgx(unique_data, extended_data, section_paths)
            # Dois níveis distintos, deliberadamente separados:
            #   has_guideline           -> o par fármaco-gene está no documento
            #   substance_has_guideline -> existe guideline para o fármaco,
            #                              mencione-a o RCM ou não
            # A diferença entre os dois é um resultado em si: mede quantos
            # fármacos com farmacogenética conhecida têm um RCM que a omite.
            has_guideline = self._has_guideline(unique_data, extended_data)
            substance_has_guideline = self._substance_has_guideline(extended_data)

            guideline_ids = self._guideline_ids_for_doc(unique_data, extended_data)
            substance_guideline_ids = self._substance_guideline_ids(extended_data)
            guideline_sources = self._guideline_sources_for_doc(extended_data)

            if not guideline_sources and guideline_ids:
                guideline_sources = "ClinPGx"

            if has_pgx:
                docs_with_pgx += 1
            else:
                docs_without_pgx += 1

            if has_pgx and has_entities:
                docs_with_pgx_entities += 1

            if has_pgx and not has_entities:
                docs_with_pgx_no_entities += 1

            if has_guideline:
                docs_with_guideline += 1
                if has_pgx:
                    docs_with_guideline_and_pgx += 1
                else:
                    docs_with_guideline_without_pgx += 1
            else:
                docs_without_guideline += 1

            if substance_has_guideline:
                docs_substance_with_guideline += 1
                if not has_guideline:
                    docs_substance_guideline_not_in_label += 1

            # --- origem das entidades: texto vs tabelas -----------------------
            origem = self._origin_block(unique_data)
            origin_rows.append({
                "document": doc_name,
                "active_substance": active_substance,
                **origem,
                "exclusivas_de_tabelas_txt": self._join_values(
                    origem["exclusivas_de_tabelas"]),
            })
            total_origin["texto_mencoes"] += origem["texto_mencoes"]
            total_origin["tabelas_mencoes"] += origem["tabelas_mencoes"]
            if origem["tabelas_mencoes"]:
                total_origin["docs_com_tabelas"] += 1
            if origem["exclusivas_de_tabelas"]:
                total_origin["docs_com_exclusivas"] += 1
                total_origin["entidades_exclusivas"] += len(origem["exclusivas_de_tabelas"])

            # --- qualidade da extração e proveniência -------------------------
            extraction_failures, non_verbatim = self._extraction_quality(unique_data)
            total_extraction_failures += extraction_failures
            total_non_verbatim += non_verbatim
            if extraction_failures:
                docs_with_extraction_failures += 1
            if non_verbatim:
                docs_with_non_verbatim += 1

            pipeline = self._pipeline_block(unique_data)
            pipeline_versions[str(pipeline.get("pipeline_version") or "desconhecida")] += 1
            llm = self._safe_dict(pipeline.get("llm"))
            if llm:
                llm_configs[
                    f"{llm.get('model')} temp={llm.get('temperature')} seed={llm.get('seed')}"
                ] += 1

            # --- cobertura de guidelines -------------------------------------
            doc_coverage = self._guideline_coverage_rows(
                doc_name, active_substance, unique_data, extended_data,
            )

            # Agregado por documento, para a folha "Documentos": une os genes
            # de todas as guidelines da substância. Responde diretamente a
            # "este RCM tem guideline mas faltam-lhe genes, e quais".
            doc_expected_genes: set[str] = set()
            doc_covered_genes: set[str] = set()
            for row in doc_coverage:
                doc_expected_genes.update(row["expected_genes"])
                doc_covered_genes.update(row["covered_genes"])
            doc_missing_genes = sorted(doc_expected_genes - doc_covered_genes)

            # Pares distintos documento × gene, sem repetição.
            #
            # A contagem ponderada soma o mesmo par uma vez por cada guideline
            # que o cobre: capecitabina–DPYD entra cinco vezes porque cinco
            # organismos publicaram sobre ele. Isso responde a "que fracção
            # das recomendações publicadas está no RCM", mas dá peso cinco a
            # um fármaco bem coberto e peso um a outro mal coberto.
            #
            # Para comparar fármacos entre si é preciso a versão sem
            # repetição: um par fármaco-gene é um facto, e não vale mais por
            # mais gente ter escrito sobre ele. Medido neste corpus: 65,52%
            # ponderado contra 49,49% em pares distintos, 16 pontos de
            # diferença.
            distinct_expected.update((doc_name, g) for g in doc_expected_genes)
            distinct_covered.update((doc_name, g) for g in doc_covered_genes)

            if doc_covered_genes and doc_missing_genes:
                docs_with_guideline_missing_genes += 1

            if doc_coverage:
                guideline_coverage_rows.extend(doc_coverage)

                for row in doc_coverage:
                    coverage_status_counts[row["status"]] += 1
                    genes_expected_total += row["n_expected"]
                    genes_covered_total += row["n_covered"]
                    # Duas contagens distintas, e confundi-las é o erro que a
                    # folha "Cobertura por gene" tinha: a coluna dizia
                    # "documentos" e somava POSIÇÕES.
                    #
                    # DPYD era esperado em 10 posições mas apenas em 2
                    # documentos (capecitabina e fluorouracilo), porque cinco
                    # organismos publicaram uma guideline DPYD para cada um.
                    # Quem lesse "DPYD: 10 documentos" concluía que dez rótulos
                    # diferentes deviam mencionar o gene, quando são dois.
                    for gene in row["expected_genes"]:
                        gene_expected_positions[gene] += 1
                        gene_expected_docs_set[gene].add(doc_name)
                    for gene in row["covered_genes"]:
                        gene_covered_positions[gene] += 1
                        gene_covered_docs_set[gene].add(doc_name)

                statuses = {row["status"] for row in doc_coverage}
                if statuses == {"coberta"}:
                    docs_fully_covered += 1
                elif statuses == {"ausente"}:
                    docs_not_covered += 1
                else:
                    docs_partially_covered += 1

            for gene in unique_entities["genes"]:
                global_unique_gene_docs[gene] += 1
                gene_mais_docs[gene].add(doc_name)

            # Alelos e diplótipos contam para o gene a que pertencem:
            # CYP2C19*2 e CYP2C19*1*2 são menções de CYP2C19. Sem isto, um
            # gene que os rótulos refiram sobretudo por alelo aparece no
            # fundo da tabela de frequências.
            for categoria in ("star_alleles", "diplotypes"):
                for entidade in unique_entities[categoria]:
                    gene = self._gene_de_entidade(entidade)
                    if gene:
                        gene_mais_docs[gene].add(doc_name)

            for gene, n in doc_mentions["genes"].items():
                gene_mais_mencoes[gene] += n
            for categoria in ("star_alleles", "diplotypes"):
                for entidade, n in doc_mentions[categoria].items():
                    gene = self._gene_de_entidade(entidade)
                    if gene:
                        gene_mais_mencoes[gene] += n

            # Um par "*x*y" é um diplótipo, mesmo que o modelo o tenha
            # devolvido na lista de alelos. Reencaminha-se aqui, na origem,
            # para que o total do resumo e a folha de frequências digam o
            # mesmo número — filtrar só na folha criava a discrepância de
            # "19 alelos únicos" no resumo contra 18 na folha.
            for allele in unique_entities["star_alleles"]:
                if self._is_diplotype(allele):
                    global_unique_diplotype_docs[allele] += 1
                else:
                    global_unique_allele_docs[allele] += 1

            for diplotype in unique_entities["diplotypes"]:
                global_unique_diplotype_docs[diplotype] += 1

            for rsid in unique_entities["rsids"]:
                global_unique_rsid_docs[rsid] += 1

            global_gene_mentions.update(doc_mentions["genes"])
            for chave, n in doc_mentions["star_alleles"].items():
                if self._is_diplotype(chave):
                    global_diplotype_mentions[chave] += n
                else:
                    global_allele_mentions[chave] += n
            global_diplotype_mentions.update(doc_mentions["diplotypes"])
            global_rsid_mentions.update(doc_mentions["rsids"])

            for gene_item in self._entities(unique_data, "genes"):
                if not isinstance(gene_item, dict):
                    continue

                gene = self._gene_value(gene_item)
                ids_block = self._safe_dict(gene_item.get("ids"))

                has_ids = any([
                    bool(gene_item.get("pharmgkb_id")),
                    bool(gene_item.get("ncbi_gene_id")),
                    bool(gene_item.get("hgnc_id")),
                    bool(gene_item.get("ensembl_id")),
                    bool(ids_block.get("pharmgkb_id")),
                    bool(ids_block.get("ncbi_gene_id")),
                    bool(ids_block.get("hgnc_id")),
                    bool(ids_block.get("ensembl_id")),
                ])

                if gene and not has_ids:
                    genes_without_tsv_docs[gene].add(doc_name)

            entity_guideline_flags = {
                "genes": set(),
                "star_alleles": set(),
                "diplotypes": set(),
                "rsids": set(),
            }

            for category, label in [
                ("genes", "gene"),
                ("star_alleles", "star_allele"),
                ("diplotypes", "diplotype"),
                ("rsids", "rsid"),
            ]:
                for item in self._entities(unique_data, category):
                    if not isinstance(item, dict):
                        continue

                    value = self._gene_value(item) if category == "genes" else self._entity_value(item)
                    has_entity_guideline = self._entity_has_guideline(item)
                    item_guideline_ids = self._entity_guideline_ids(item)

                    if has_entity_guideline and value:
                        entity_guideline_flags[category].add(value)

                    entidades_clinpgx_rows.append({
                        "document": doc_name,
                        "active_substance": active_substance,
                        "entity": value,
                        "type": label,
                        "total_mentions_real": doc_mentions[category].get(value, 0),
                        "has_clinpgx": has_entity_guideline,
                        "guideline_ids": self._join_values(item_guideline_ids),
                    })

            for gene in entity_guideline_flags["genes"]:
                total_gene_mentions_with_guideline += doc_mentions["genes"].get(gene, 0)

            for category in ["genes", "star_alleles", "diplotypes", "rsids"]:
                for entity in entity_guideline_flags[category]:
                    total_entity_mentions_with_guideline += doc_mentions[category].get(entity, 0)

            extended_norm_rows = self._extended_normalization_rows(
                extended_data=extended_data,
                doc_name=doc_name,
                active_substance=active_substance,
            )

            combined_norm_rows = []
            combined_norm_rows.extend(extended_norm_rows)
            combined_norm_rows.extend(doc_normalization_rows)

            combined_norm_rows = self._aggregate_normalization_rows_if_needed(
                unique_data=unique_data,
                extended_data=extended_data,
                doc_name=doc_name,
                active_substance=active_substance,
                existing_rows=combined_norm_rows,
            )

            normalizacao_rows.extend(combined_norm_rows)

            total_model_normalizations += self._to_int(
                counts.get("genes_normalizados_nome_extenso_para_simbolo_total")
            )
            total_technical_normalizations += self._to_int(
                counts.get("normalizacoes_tecnicas_total")
            )
            total_incomplete_star_alleles += self._to_int(
                counts.get("star_alleles_incompletos_total")
            )

            gene_mentions_real = sum(doc_mentions["genes"].values())
            star_mentions_real = sum(doc_mentions["star_alleles"].values())
            diplotype_mentions_real = sum(doc_mentions["diplotypes"].values())
            rsid_mentions_real = sum(doc_mentions["rsids"].values())

            documents.append({
                "document": doc_name,
                "medicamento": medicamento,
                "active_substance": active_substance,
                "active_substance_aliases": aliases,
                "grupo_farmacoterapeutico": grupo,
                "grupo_farmacoterapeutico_codigo": grupo_codigo,
                "atc": atc,
                "mesh_id": ids["mesh_id"],
                "drugbank_id": ids["drugbank_id"],
                "chebi_id": ids["chebi_id"],
                "umls_cui": ids["umls_cui"],
                "substance_has_guideline": substance_has_guideline,
                "substance_guideline_ids": self._join_values(substance_guideline_ids),
                "guideline_genes_expected": self._join_values(doc_expected_genes),
                "guideline_genes_covered": self._join_values(doc_covered_genes),
                "guideline_genes_missing": self._join_values(doc_missing_genes),
                "guideline_gene_coverage_pct": self._pct(
                    len(doc_covered_genes), len(doc_expected_genes)
                ),
                "extraction_failures": extraction_failures,
                "non_verbatim_excerpts": non_verbatim,
                "has_pgx": has_pgx,
                "has_entities": has_entities,
                "has_guideline": has_guideline,
                "guideline_ids": self._join_values(guideline_ids),
                "guideline_sources": guideline_sources,
                "unique_genes": len(unique_entities["genes"]),
                "gene_mentions": gene_mentions_real,
                "unique_star_alleles": len(unique_entities["star_alleles"]),
                "star_allele_mentions": star_mentions_real,
                "unique_diplotypes": len(unique_entities["diplotypes"]),
                "diplotype_mentions": diplotype_mentions_real,
                "unique_rsids": len(unique_entities["rsids"]),
                "rsid_mentions": rsid_mentions_real,
                "model_gene_normalizations": self._to_int(
                    counts.get("genes_normalizados_nome_extenso_para_simbolo_total")
                ),
                "technical_normalizations": self._to_int(
                    counts.get("normalizacoes_tecnicas_total")
                ),
                "incomplete_star_alleles": self._to_int(
                    counts.get("star_alleles_incompletos_total")
                ),
                "status": "ok",
            })

            farmaco_entidades.append({
                "document": doc_name,
                "medicamento": medicamento,
                "active_substance": active_substance,
                "gene_occurrences": self._format_counter(doc_mentions["genes"]),
                "star_allele_occurrences": self._format_counter(doc_mentions["star_alleles"]),
                "diplotype_occurrences": self._format_counter(doc_mentions["diplotypes"]),
                "rsid_occurrences": self._format_counter(doc_mentions["rsids"]),
                "has_guideline": has_guideline,
                "guideline_ids": self._join_values(guideline_ids),
                "guideline_sources": guideline_sources,
            })

            # Agrupar pelo código hierárquico do Infarmed, não pelo rótulo
            # textual: o mesmo grupo aparece escrito de dezenas de maneiras
            # (maiúsculas, hífens, espaçamento, truncagens), o que fragmentava
            # 121 grupos reais em 472 linhas distintas na folha do Excel.
            group_key = grupo_codigo or grupo or "(sem grupo)"

            # Guarda o rótulo mais completo visto para este código, para o
            # apresentar no Excel em vez do código nu.
            label = grupo or "(sem grupo)"
            if len(label) > len(group_labels.get(group_key, "")):
                group_labels[group_key] = label

            group_stats[group_key]["documents"].add(doc_name)

            if has_pgx:
                group_stats[group_key]["documents_with_pgx"].add(doc_name)

            if has_pgx and has_entities:
                group_stats[group_key]["documents_with_pgx_entities"].add(doc_name)

            if has_guideline:
                group_stats[group_key]["documents_with_guideline"].add(doc_name)

            group_stats[group_key]["gene_mentions"] += gene_mentions_real
            group_stats[group_key]["unique_genes"].update(unique_entities["genes"])

        total_gene_mentions = sum(global_gene_mentions.values())
        total_allele_mentions = sum(global_allele_mentions.values())
        total_diplotype_mentions = sum(global_diplotype_mentions.values())
        total_rsid_mentions = sum(global_rsid_mentions.values())

        total_unique_genes = len(global_unique_gene_docs)
        total_unique_alleles = len(global_unique_allele_docs)
        total_unique_diplotypes = len(global_unique_diplotype_docs)
        total_unique_rsids = len(global_unique_rsid_docs)

        all_entity_mentions = (
            total_gene_mentions
            + total_allele_mentions
            + total_diplotype_mentions
            + total_rsid_mentions
        )

        section_summary = []
        for section, stats in section_stats.items():
            section_summary.append({
                "section": section,
                "documents_count": len(stats["documents"]),
                "documents_with_pgx_count": len(stats["documents_with_pgx"]),
                "gene_mentions": stats["gene_mentions"],
                "star_allele_mentions": stats["star_allele_mentions"],
                "diplotype_mentions": stats["diplotype_mentions"],
                "rsid_mentions": stats["rsid_mentions"],
                "total_mentions": stats["total_mentions"],
                "percentage_of_all_entity_mentions": self._pct(
                    stats["total_mentions"],
                    all_entity_mentions,
                ),
            })

        section_summary.sort(key=lambda x: x["total_mentions"], reverse=True)

        group_summary = []
        for group, stats in group_stats.items():
            group_summary.append({
                "codigo_grupo": group if group != "(sem grupo)" else "",
                "grupo_farmacoterapeutico": group_labels.get(group, group),
                "documents_count": len(stats["documents"]),
                "documents_with_pgx": len(stats["documents_with_pgx"]),
                "documents_with_pgx_entities": len(stats["documents_with_pgx_entities"]),
                "documents_with_guideline": len(stats["documents_with_guideline"]),
                "gene_mentions": stats["gene_mentions"],
                "unique_genes": len(stats["unique_genes"]),
                "percentage_pgx": self._pct(
                    len(stats["documents_with_pgx"]),
                    len(stats["documents"]),
                ),
            })

        group_summary.sort(
            key=lambda x: (x["documents_with_pgx"], x["gene_mentions"]),
            reverse=True,
        )

        result = {
            "summary": {
                "total_documents_processed": total_docs,
                "documents_with_document_unique": docs_with_unique,
                # Excluded from every denominator above: no readable text, so
                # no chance of entities. Reported so the exclusion is auditable.
                "documents_unusable_excluded": len(self.unusable_dirs),
                "documents_unusable_names": [p.name for p in self.unusable_dirs],
                "documents_unusable_detail": self._unusable_rows(),
                "documents_with_pgx": docs_with_pgx,
                "documents_without_pgx": docs_without_pgx,
                "documents_with_pgx_and_entities": docs_with_pgx_entities,
                "documents_with_pgx_without_entities": docs_with_pgx_no_entities,
                "documents_with_guideline": docs_with_guideline,
                "documents_without_guideline": docs_without_guideline,
                "documents_with_guideline_and_pgx": docs_with_guideline_and_pgx,
                "documents_with_guideline_without_pgx": docs_with_guideline_without_pgx,
                "documents_substance_with_guideline": docs_substance_with_guideline,
                "documents_substance_guideline_not_in_label": docs_substance_guideline_not_in_label,
                # cobertura de guidelines
                "guideline_pairs_total": sum(coverage_status_counts.values()),
                "guideline_pairs_covered": coverage_status_counts.get("coberta", 0),
                "guideline_pairs_partial": coverage_status_counts.get("parcial", 0),
                "guideline_pairs_absent": coverage_status_counts.get("ausente", 0),
                "guideline_genes_expected": genes_expected_total,
                "guideline_genes_covered": genes_covered_total,
                "percentage_guideline_genes_covered": self._pct(
                    genes_covered_total, genes_expected_total
                ),
                # Sem repetição: cada par fármaco-gene conta uma vez.
                "distinct_drug_gene_expected": len(distinct_expected),
                "distinct_drug_gene_covered": len(distinct_covered),
                "percentage_distinct_drug_gene_covered": self._pct(
                    len(distinct_covered), len(distinct_expected)
                ),
                "documents_fully_covered": docs_fully_covered,
                "documents_partially_covered": docs_partially_covered,
                "documents_not_covered": docs_not_covered,
                "documents_with_guideline_missing_genes": docs_with_guideline_missing_genes,
                # qualidade e proveniência
                "total_extraction_failures": total_extraction_failures,
                "documents_with_extraction_failures": docs_with_extraction_failures,
                "total_non_verbatim_excerpts": total_non_verbatim,
                "documents_with_non_verbatim_excerpts": docs_with_non_verbatim,
                "pipeline_versions": dict(pipeline_versions),
                "llm_configs": dict(llm_configs),
                "percentage_with_pgx": self._pct(docs_with_pgx, docs_with_unique),
                "percentage_with_pgx_and_entities": self._pct(docs_with_pgx_entities, docs_with_unique),
                "percentage_with_pgx_without_entities": self._pct(docs_with_pgx_no_entities, docs_with_unique),
                "percentage_with_guideline": self._pct(docs_with_guideline, docs_with_unique),
                "total_unique_genes": total_unique_genes,
                "total_gene_mentions": total_gene_mentions,
                "total_unique_star_alleles": total_unique_alleles,
                "total_star_allele_mentions": total_allele_mentions,
                "total_unique_diplotypes": total_unique_diplotypes,
                "total_diplotype_mentions": total_diplotype_mentions,
                "total_unique_rsids": total_unique_rsids,
                "total_rsid_mentions": total_rsid_mentions,
                "total_entity_mentions": all_entity_mentions,
                "total_model_gene_normalizations": total_model_normalizations,
                "total_technical_normalizations": total_technical_normalizations,
                "total_incomplete_star_alleles": total_incomplete_star_alleles,
                "total_gene_mentions_with_guideline": total_gene_mentions_with_guideline,
                "total_entity_mentions_with_guideline": total_entity_mentions_with_guideline,
            },
            "documents": documents,
            "origem_entidades": origin_rows,
            "origem_totais": dict(total_origin),
            "guideline_coverage": guideline_coverage_rows,
            "guideline_gene_coverage": [
                {
                    "gene": gene,
                    # Documentos distintos: quantos RCMs deviam mencionar este
                    # gene, e quantos o fazem. É a leitura intuitiva.
                    "documents_expected": len(gene_expected_docs_set[gene]),
                    "documents_covered": len(gene_covered_docs_set[gene]),
                    "percentage_documents_covered": self._pct(
                        len(gene_covered_docs_set[gene]),
                        len(gene_expected_docs_set[gene]),
                    ),
                    # Posições: o mesmo par conta uma vez por guideline. É o
                    # que soma 174 e alimenta a percentagem ponderada.
                    "positions_expected": n_pos,
                    "positions_covered": gene_covered_positions.get(gene, 0),
                    "percentage_positions_covered": self._pct(
                        gene_covered_positions.get(gene, 0), n_pos
                    ),
                }
                for gene, n_pos in gene_expected_positions.most_common()
            ],
            "farmaco_entidades": farmaco_entidades,
            "section_summary": section_summary,
            "section_rows": section_rows,
            "group_summary": group_summary,
            "normalizacoes": normalizacao_rows,
            "entidades_clinpgx": entidades_clinpgx_rows,
            "genes_agregados": [
                {
                    "gene": gene,
                    "mencoes_simbolo": global_gene_mentions.get(gene, 0),
                    "mencoes_via_alelo_diplotipo":
                        n - global_gene_mentions.get(gene, 0),
                    "mencoes_total": n,
                    "documentos_uniao": len(gene_mais_docs.get(gene, set())),
                }
                for gene, n in gene_mais_mencoes.most_common()
            ],
            "genes_frequentes": self._build_frequency_rows(
                global_unique_gene_docs,
                global_gene_mentions,
                "gene",
                docs_with_unique,
                docs_with_pgx,
            ),
            "alelos_frequentes": self._build_frequency_rows(
                global_unique_allele_docs,
                global_allele_mentions,
                "star_allele",
                docs_with_unique,
                docs_with_pgx,
            ),
            "diplotipos_frequentes": self._build_frequency_rows(
                global_unique_diplotype_docs,
                global_diplotype_mentions,
                "diplotype",
                docs_with_unique,
                docs_with_pgx,
            ),
            "rsids_frequentes": self._build_frequency_rows(
                global_unique_rsid_docs,
                global_rsid_mentions,
                "rsid",
                docs_with_unique,
                docs_with_pgx,
            ),
            "genes_sem_tsv": [
                {
                    "gene": gene,
                    "documents_count": len(docs),
                    "documents": sorted(docs),
                }
                for gene, docs in sorted(genes_without_tsv_docs.items())
            ],
        }

        (self.outputs_root / self.JSON_OUTPUT).write_text(
            json.dumps(result, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        self._export_xlsx(result)

        return result

    @staticmethod
    def _is_diplotype(entidade: str) -> bool:
        """Um par de alelos, e não um alelo.

        "CYP2C9*2*3" chegava à lista de alelos porque a normalização do
        siponimod devolveu "*2*3" e a entidade ficou marcada `star_allele`.
        Ficava contado duas vezes, em duas categorias, e a folha de alelos
        passava a incluir um par.

        A forma distingue-os sem ambiguidade: um alelo tem um sufixo estrela
        (CYP2C9*3, HLA-B*1502), um diplótipo tem dois.
        """
        return str(entidade).count("*") > 1

    def _build_frequency_rows(
        self,
        doc_counter: Counter,
        mention_counter: Counter,
        entity_key: str,
        total_docs: int,
        total_pgx_docs: int,
    ) -> list[dict]:
        entities = set(doc_counter.keys()) | set(mention_counter.keys())
        rows = []

        for entity in entities:
            rows.append({
                entity_key: entity,
                "documents_count": doc_counter.get(entity, 0),
                "total_mentions": mention_counter.get(entity, 0),
                "percentage_total_documents": self._pct(
                    doc_counter.get(entity, 0),
                    total_docs,
                ),
                "percentage_pgx_documents": self._pct(
                    doc_counter.get(entity, 0),
                    total_pgx_docs,
                ),
            })

        rows.sort(
            key=lambda x: (x["total_mentions"], x["documents_count"]),
            reverse=True,
        )

        return rows

    # =====================================================
    # XLSX
    # =====================================================
    def _export_xlsx(self, result: dict) -> Path:
        wb = Workbook()

        # Ordem por tema, do geral para o detalhe.
        #
        # A ordem anterior misturava níveis: a cobertura resumo vinha antes da
        # origem das entidades, e o detalhe da cobertura depois — obrigando a
        # saltar para trás e para a frente para seguir o mesmo raciocínio.
        # Agora cada bloco fecha antes do seguinte começar, e dentro de cada
        # bloco vai-se do agregado para a linha individual.

        # A. Panorama
        self._write_resumo(wb, result)
        self._write_documents(wb, result)

        # B. O que foi encontrado
        self._write_pgx_com_vs_sem_entidades(wb, result)
        self._write_ocorrencias_vs_unicos(wb, result)
        self._write_frequency_sheet(wb, "Genes frequentes", result["genes_frequentes"], "gene")
        self._write_genes_agregados(wb, result)
        self._write_frequency_sheet(wb, "Alelos frequentes", result["alelos_frequentes"], "star_allele")
        self._write_frequency_sheet(wb, "Diplótipos frequentes", result["diplotipos_frequentes"], "diplotype")
        self._write_frequency_sheet(wb, "rsIDs frequentes", result["rsids_frequentes"], "rsid")
        self._write_farmaco_entidades(wb, result)

        # C. Cobertura de guidelines — resumo, depois detalhe, depois por gene
        self._write_guidelines_coverage(wb, result)
        self._write_guideline_coverage_detail(wb, result)
        self._write_guideline_gene_coverage(wb, result)
        self._write_entidades_clinpgx(wb, result)

        # D. Onde no documento
        self._write_section_summary(wb, result)
        self._write_section_rows(wb, result)
        self._write_origem_entidades(wb, result)
        self._write_group_summary(wb, result)

        # E. Qualidade e diagnóstico
        self._write_normalizacoes(wb, result)
        self._write_formas_entidades(wb, result)
        self._write_documentos_excluidos(wb, result)
        self._write_genes_sem_tsv(wb, result)

        self._format_workbook(wb)

        output_path = self.outputs_root / self.XLSX_OUTPUT
        wb.save(output_path)

        return output_path

    def _write_resumo(self, wb: Workbook, result: dict) -> None:
        """Resumo em blocos temáticos, não em lista corrida.

        Uma coluna única de 35 métricas lê-se de cima para baixo e convida a
        comparar números que não são comparáveis entre si: o denominador muda
        de bloco para bloco. "Documentos com PGx" é sobre 64 documentos,
        "Pares cobertos" é sobre 145 pares, "Genes cobertos" é sobre 174
        posições de gene. Postos em fila, os três parecem a mesma coisa.

        Cada bloco leva um cabeçalho e, onde é preciso, o denominador escrito
        na própria linha.
        """
        ws = wb.active
        ws.title = "Resumo"

        s = result["summary"]
        n_docs = s["total_documents_processed"]

        def pct(parte, total):
            return round(100 * parte / total, 2) if total else 0.0

        blocos = [
            ("1. CORPUS", [
                ("Documentos analisados", n_docs, ""),
                ("Documentos com Document_Unique",
                 s["documents_with_document_unique"], ""),
                ("Excluídos por não terem texto utilizável",
                 s.get("documents_unusable_excluded", 0),
                 "digitalizações e ficheiros com texto corrompido; "
                 "fora de todos os denominadores"),
            ]),
            ("2. CONTEÚDO FARMACOGENÓMICO   (denominador: documentos)", [
                ("Documentos com conteúdo PGx", s["documents_with_pgx"],
                 f"{pct(s['documents_with_pgx'], n_docs)}% de {n_docs}"),
                ("   dos quais com entidades extraídas",
                 s["documents_with_pgx_and_entities"],
                 f"{pct(s['documents_with_pgx_and_entities'], n_docs)}% de {n_docs}"),
                ("   dos quais sem nenhuma entidade",
                 s["documents_with_pgx_without_entities"],
                 "o RCM refere genética sem nomear gene, alelo ou diplótipo"),
                ("Documentos sem conteúdo PGx", s["documents_without_pgx"], ""),
            ]),
            ("3. ENTIDADES EXTRAÍDAS   (denominador: menções)", [
                ("Genes únicos", s["total_unique_genes"], ""),
                ("Menções de genes", s["total_gene_mentions"],
                 f"das quais {s['total_gene_mentions_with_guideline']} "
                 f"em genes com guideline do fármaco"),
                ("Alelos únicos", s["total_unique_star_alleles"], ""),
                ("Menções de alelos", s["total_star_allele_mentions"], ""),
                ("Diplótipos únicos", s["total_unique_diplotypes"], ""),
                ("Menções de diplótipos", s["total_diplotype_mentions"], ""),
                ("rsIDs únicos", s["total_unique_rsids"], ""),
                ("Menções de rsIDs", s["total_rsid_mentions"], ""),
                ("Menções de entidades (total)", s["total_entity_mentions"],
                 f"das quais {s['total_entity_mentions_with_guideline']} "
                 f"com guideline"),
            ]),
            ("4. GUIDELINES — o fármaco tem?   (denominador: documentos)", [
                ("Substância com guideline no ClinPGx",
                 s["documents_substance_with_guideline"],
                 "propriedade do fármaco; independente do que o RCM diz"),
                ("   e o RCM menciona pelo menos um gene",
                 s["documents_with_guideline"],
                 f"{pct(s['documents_with_guideline'], n_docs)}% de {n_docs}"),
                ("   mas o RCM não menciona gene nenhum",
                 s["documents_substance_guideline_not_in_label"],
                 "lacuna regulamentar"),
                ("Substância sem guideline no ClinPGx",
                 n_docs - s["documents_substance_with_guideline"], ""),
            ]),
            ("5. COBERTURA — pares documento × guideline   (denominador: pares)", [
                ("Pares avaliados", s["guideline_pairs_total"],
                 "uma linha por cada guideline de cada substância"),
                ("   cobertos (todos os genes mencionados)",
                 s["guideline_pairs_covered"],
                 f"{pct(s['guideline_pairs_covered'], s['guideline_pairs_total'])}%"),
                ("   parciais", s["guideline_pairs_partial"],
                 "só possível em guidelines de mais de um gene"),
                ("   ausentes", s["guideline_pairs_absent"],
                 f"{pct(s['guideline_pairs_absent'], s['guideline_pairs_total'])}%"),
            ]),
            ("6. COBERTURA — ao nível do gene   (denominador: posições de gene)", [
                ("Posições de gene esperadas",
                 s["guideline_genes_expected"],
                 "COM repetição: o mesmo par fármaco-gene conta uma vez por "
                 "cada guideline que o cobre"),
                ("Posições cobertas", s["guideline_genes_covered"], ""),
                ("% cobertura (ponderada por guideline)",
                 s["percentage_guideline_genes_covered"],
                 "responde a: que fracção das recomendações publicadas está "
                 "no RCM?"),
                ("Pares distintos fármaco-gene esperados",
                 s.get("distinct_drug_gene_expected", ""),
                 "SEM repetição: cada par conta uma vez"),
                ("Pares distintos cobertos",
                 s.get("distinct_drug_gene_covered", ""), ""),
                ("% cobertura (pares distintos)",
                 s.get("percentage_distinct_drug_gene_covered", ""),
                 "responde a: que fracção do conhecimento fármaco-gene está "
                 "no RCM? — a métrica a usar para comparar fármacos"),
            ]),
            ("7. COBERTURA — ao nível do documento   (denominador: documentos com guideline)", [
                ("Totalmente cobertos", s["documents_fully_covered"], ""),
                ("Parcialmente cobertos", s["documents_partially_covered"], ""),
                ("Sem cobertura", s["documents_not_covered"], ""),
            ]),
            ("8. QUALIDADE DA EXTRAÇÃO", [
                ("Falhas de extração (resposta não utilizável)",
                 s["total_extraction_failures"],
                 "0 significa que o modelo devolveu sempre JSON válido"),
                ("Documentos com pelo menos uma falha",
                 s["documents_with_extraction_failures"], ""),
                ("Excertos não-verbatim", s["total_non_verbatim_excerpts"],
                 "ver folha Normalizações; secções longas são partidas em "
                 "blocos e isto sobrestima o problema"),
                ("Documentos com excertos não-verbatim",
                 s["documents_with_non_verbatim_excerpts"], ""),
                ("Normalizações nome extenso → símbolo",
                 s["total_model_gene_normalizations"],
                 "eventos, não menções"),
                ("Normalizações técnicas",
                 s["total_technical_normalizations"], "eventos"),
                ("Alelos incompletos", s["total_incomplete_star_alleles"],
                 "eventos"),
            ]),
            ("9. PROVENIÊNCIA", [
                ("Versões do pipeline", " ; ".join(
                    f"{k}={v}" for k, v in sorted(s["pipeline_versions"].items())),
                 "mais do que uma versão significa corpus não comparável"),
                ("Configurações do modelo", " ; ".join(
                    f"{k} ({v} docs)" for k, v in sorted(s["llm_configs"].items())),
                 ""),
            ]),
        ]

        for titulo, linhas in blocos:
            ws.append([titulo])
            ws.append(["Métrica", "Valor", "Nota"])
            for nome, valor, nota in linhas:
                ws.append([nome, valor, nota])
            ws.append([])

    def _write_documents(self, wb: Workbook, result: dict) -> None:
        ws = wb.create_sheet("Documentos")

        ws.append([
            "Documento",
            "Medicamento",
            "Substância ativa",
            "Aliases",
            "Código grupo",
            "Grupo farmacoterapêutico",
            "ATC",
            "MeSH ID",
            "DrugBank ID",
            "ChEBI ID",
            "UMLS CUI",
            "Tem PGx",
            "Tem entidades",
            "Tem guideline (par fármaco-gene)",
            "Guideline IDs",
            "Substância tem guideline",
            "Guideline IDs da substância",
            "Genes de guideline esperados",
            "Genes de guideline mencionados",
            "Genes de guideline EM FALTA",
            "% cobertura de genes",
            "Falhas de extração",
            "Excertos não-verbatim",
            "Fontes guideline",
            "Genes únicos",
            "Menções genes reais",
            "Alelos únicos",
            "Menções alelos reais",
            "Diplótipos únicos",
            "Menções diplótipos reais",
            "rsIDs únicos",
            "Menções rsIDs reais",
            "Normalizações nome extenso → símbolo",
            "Normalizações técnicas",
            "Star alleles incompletos",
            "Estado",
        ])

        for item in result["documents"]:
            ws.append([
                item.get("document", ""),
                item.get("medicamento", ""),
                item.get("active_substance", ""),
                " ; ".join(item.get("active_substance_aliases") or []),
                item.get("grupo_farmacoterapeutico_codigo", ""),
                item.get("grupo_farmacoterapeutico", ""),
                item.get("atc", ""),
                item.get("mesh_id", ""),
                item.get("drugbank_id", ""),
                item.get("chebi_id", ""),
                item.get("umls_cui", ""),
                "Sim" if item.get("has_pgx") else "Não",
                "Sim" if item.get("has_entities") else "Não",
                "Sim" if item.get("has_guideline") else "Não",
                item.get("guideline_ids", ""),
                "Sim" if item.get("substance_has_guideline") else "Não",
                item.get("substance_guideline_ids", ""),
                item.get("guideline_genes_expected", ""),
                item.get("guideline_genes_covered", ""),
                item.get("guideline_genes_missing", ""),
                item.get("guideline_gene_coverage_pct", 0),
                item.get("extraction_failures", 0),
                item.get("non_verbatim_excerpts", 0),
                item.get("guideline_sources", ""),
                item.get("unique_genes", 0),
                item.get("gene_mentions", 0),
                item.get("unique_star_alleles", 0),
                item.get("star_allele_mentions", 0),
                item.get("unique_diplotypes", 0),
                item.get("diplotype_mentions", 0),
                item.get("unique_rsids", 0),
                item.get("rsid_mentions", 0),
                item.get("model_gene_normalizations", 0),
                item.get("technical_normalizations", 0),
                item.get("incomplete_star_alleles", 0),
                item.get("status", ""),
            ])

    def _write_pgx_com_vs_sem_entidades(self, wb: Workbook, result: dict) -> None:
        ws = wb.create_sheet("PGx com vs sem entidades")
        s = result["summary"]

        ws.append(["Categoria", "Valor", "Percentagem sobre Document_Unique"])
        ws.append([
            "Documentos com PGx e entidades",
            s["documents_with_pgx_and_entities"],
            s["percentage_with_pgx_and_entities"],
        ])
        ws.append([
            "Documentos com PGx sem entidades",
            s["documents_with_pgx_without_entities"],
            s["percentage_with_pgx_without_entities"],
        ])

    def _write_ocorrencias_vs_unicos(self, wb: Workbook, result: dict) -> None:
        ws = wb.create_sheet("Ocorrências vs Únicos")
        s = result["summary"]

        ws.append(["Tipo", "Entidades únicas", "Menções totais reais"])
        ws.append(["Genes", s["total_unique_genes"], s["total_gene_mentions"]])
        ws.append(["Alelos", s["total_unique_star_alleles"], s["total_star_allele_mentions"]])
        ws.append(["Diplótipos", s["total_unique_diplotypes"], s["total_diplotype_mentions"]])
        ws.append(["rsIDs", s["total_unique_rsids"], s["total_rsid_mentions"]])

    def _write_frequency_sheet(
        self,
        wb: Workbook,
        title: str,
        rows: list[dict],
        key: str,
    ) -> None:
        ws = wb.create_sheet(title)

        ws.append([
            "Entidade",
            "Nº documentos onde aparece",
            "Nº menções totais reais",
            "% total documentos",
            "% documentos PGx",
        ])

        for item in rows:
            ws.append([
                item.get(key, ""),
                item.get("documents_count", 0),
                item.get("total_mentions", 0),
                item.get("percentage_total_documents", 0),
                item.get("percentage_pgx_documents", 0),
            ])

    def _write_section_summary(self, wb: Workbook, result: dict) -> None:
        ws = wb.create_sheet("Entidades por secção")

        ws.append([
            "Secção",
            "Nº documentos",
            "Nº documentos PGx",
            "Menções genes",
            "Menções alelos",
            "Menções diplótipos",
            "Menções rsIDs",
            "Total menções",
            "% das menções totais",
        ])

        for item in result["section_summary"]:
            ws.append([
                item.get("section", ""),
                item.get("documents_count", 0),
                item.get("documents_with_pgx_count", 0),
                item.get("gene_mentions", 0),
                item.get("star_allele_mentions", 0),
                item.get("diplotype_mentions", 0),
                item.get("rsid_mentions", 0),
                item.get("total_mentions", 0),
                item.get("percentage_of_all_entity_mentions", 0),
            ])

    def _write_section_rows(self, wb: Workbook, result: dict) -> None:
        ws = wb.create_sheet("Secções detalhe")

        ws.append([
            "Documento",
            "Secção",
            "Menções genes",
            "Menções alelos",
            "Menções diplótipos",
            "Menções rsIDs",
            "Total menções",
            "Genes",
            "Alelos",
            "Diplótipos",
            "rsIDs",
        ])

        for item in result["section_rows"]:
            ws.append([
                item.get("document", ""),
                item.get("section", ""),
                item.get("gene_mentions", 0),
                item.get("star_allele_mentions", 0),
                item.get("diplotype_mentions", 0),
                item.get("rsid_mentions", 0),
                item.get("total_mentions", 0),
                item.get("genes", ""),
                item.get("star_alleles", ""),
                item.get("diplotypes", ""),
                item.get("rsids", ""),
            ])

    def _write_group_summary(self, wb: Workbook, result: dict) -> None:
        ws = wb.create_sheet("Grupo farmacoterapêutico")

        ws.append([
            "Código",
            "Grupo farmacoterapêutico",
            "Nº documentos",
            "Nº documentos com PGx",
            "Nº documentos PGx com entidades",
            "Nº documentos com guideline",
            "Menções genes",
            "Genes únicos",
            "% PGx",
        ])

        for item in result["group_summary"]:
            ws.append([
                item.get("codigo_grupo", ""),
                item.get("grupo_farmacoterapeutico", ""),
                item.get("documents_count", 0),
                item.get("documents_with_pgx", 0),
                item.get("documents_with_pgx_entities", 0),
                item.get("documents_with_guideline", 0),
                item.get("gene_mentions", 0),
                item.get("unique_genes", 0),
                item.get("percentage_pgx", 0),
            ])

    def _write_farmaco_entidades(self, wb: Workbook, result: dict) -> None:
        ws = wb.create_sheet("Fármaco ↔ Entidades")

        ws.append([
            "Documento",
            "Medicamento",
            "Substância ativa",
            "Ocorrência genes reais",
            "Ocorrência alelos reais",
            "Ocorrência diplótipos reais",
            "Ocorrência rsIDs reais",
            "Contém guideline?",
            # Estes IDs são das guidelines do PAR fármaco-gene observado, não
            # de todas as guidelines da substância. O acenocumarol tem três
            # guidelines mas aqui aparecem duas: falta a de VKORC1, gene que o
            # RCM não menciona. O nome antigo ("Guideline IDs") não dizia qual
            # das duas coisas era, e as duas existem no ficheiro.
            "Guideline IDs (do par fármaco-gene observado)",
            "Fontes guideline",
        ])

        for item in result["farmaco_entidades"]:
            ws.append([
                item.get("document", ""),
                item.get("medicamento", ""),
                item.get("active_substance", ""),
                item.get("gene_occurrences", ""),
                item.get("star_allele_occurrences", ""),
                item.get("diplotype_occurrences", ""),
                item.get("rsid_occurrences", ""),
                "Sim" if item.get("has_guideline") else "Não",
                item.get("guideline_ids", ""),
                item.get("guideline_sources", ""),
            ])

    def _write_guidelines_coverage(self, wb: Workbook, result: dict) -> None:
        ws = wb.create_sheet("Cobertura resumo")
        s = result["summary"]

        ws.append(["Métrica", "Definição", "Valor"])

        ws.append([
            "Documentos com guideline (par fármaco-gene)",
            "Entidade PGx detetada no RCM que consta de uma guideline da substância",
            s["documents_with_guideline"],
        ])
        ws.append([
            "Documentos sem guideline",
            "",
            s["documents_without_guideline"],
        ])
        ws.append([
            "Documentos com guideline e PGx",
            "",
            s["documents_with_guideline_and_pgx"],
        ])
        ws.append([
            "Documentos com guideline sem PGx",
            "Deve ser 0 — com a definição por par, uma guideline exige entidade detetada",
            s["documents_with_guideline_without_pgx"],
        ])
        ws.append([
            "% documentos com guideline",
            "",
            s["percentage_with_guideline"],
        ])

        ws.append([])
        ws.append([
            "Substância tem guideline no ClinPGx",
            "Propriedade do fármaco; independente do RCM a mencionar",
            s["documents_substance_with_guideline"],
        ])
        ws.append([
            "…mas o RCM não menciona o gene",
            "Lacuna regulamentar: fármaco com PGx conhecida cujo RCM a omite",
            s["documents_substance_guideline_not_in_label"],
        ])
        ws.append([
            "% de omissão",
            "sobre as substâncias com guideline",
            self._pct(
                s["documents_substance_guideline_not_in_label"],
                s["documents_substance_with_guideline"],
            ),
        ])

        ws.append([])
        ws.append(["COBERTURA — pares documento × guideline", "", ""])
        ws.append([
            "Pares avaliados",
            "Cada guideline da substância, por documento",
            s["guideline_pairs_total"],
        ])
        ws.append([
            "Coberta",
            "O RCM menciona TODOS os genes da guideline",
            s["guideline_pairs_covered"],
        ])
        ws.append([
            "Parcial",
            "Menciona alguns; só possível em guidelines multi-gene",
            s["guideline_pairs_partial"],
        ])
        ws.append([
            "Ausente",
            "Não menciona nenhum dos genes",
            s["guideline_pairs_absent"],
        ])

        ws.append([])
        ws.append(["COBERTURA — a nível do gene", "", ""])
        ws.append([
            "Genes esperados (com repetição)",
            "Soma dos genes exigidos em todos os pares",
            s["guideline_genes_expected"],
        ])
        ws.append([
            "Genes mencionados",
            "",
            s["guideline_genes_covered"],
        ])
        ws.append([
            "% cobertura de genes de guideline",
            "A métrica principal: quanto os RCMs cobrem as guidelines",
            s["percentage_guideline_genes_covered"],
        ])

        ws.append([])
        ws.append(["COBERTURA — a nível do documento", "", ""])
        ws.append([
            "Documentos totalmente cobertos",
            "Todas as guidelines da substância cobertas",
            s["documents_fully_covered"],
        ])
        ws.append([
            "Documentos parcialmente cobertos",
            "Algumas guidelines cobertas, outras não",
            s["documents_partially_covered"],
        ])
        ws.append([
            "Documentos sem cobertura",
            "Substância tem guideline; RCM não menciona gene nenhum",
            s["documents_not_covered"],
        ])
        ws.append([
            "Documentos com guideline MAS com genes em falta",
            "Menciona pelo menos um gene e omite pelo menos um — ver coluna "
            "'Genes de guideline EM FALTA' na folha Documentos",
            s["documents_with_guideline_missing_genes"],
        ])

    def _write_origem_entidades(self, wb: Workbook, result: dict) -> None:
        """Uma folha só: texto, tabelas e total, lado a lado.

        As tabelas dos RCMs eram removidas do texto e nunca avaliadas. Esta
        folha mede o que isso custava — a coluna que interessa é
        "Entidades SÓ em tabelas": entidades que não existem em mais lado
        nenhum do documento.
        """
        # Sem ":" no nome — o Excel proíbe : \ / ? * [ ] em títulos de folha.
        ws = wb.create_sheet("Origem das entidades")
        totais = result.get("origem_totais") or {}
        linhas = result.get("origem_entidades") or []

        ws.append(["RESUMO", "", ""])
        ws.append(["Menções vindas de texto corrido", "", totais.get("texto_mencoes", 0)])
        ws.append(["Menções vindas de tabelas", "", totais.get("tabelas_mencoes", 0)])
        ws.append([
            "% das menções que vêm de tabelas", "",
            self._pct(totais.get("tabelas_mencoes", 0),
                      totais.get("texto_mencoes", 0) + totais.get("tabelas_mencoes", 0)),
        ])
        ws.append(["Documentos com entidades em tabelas", "",
                   totais.get("docs_com_tabelas", 0)])
        ws.append([
            "Documentos com entidades SÓ em tabelas",
            "seriam perdidas se as tabelas não fossem avaliadas",
            totais.get("docs_com_exclusivas", 0),
        ])
        ws.append(["Entidades exclusivas de tabelas (total)", "",
                   totais.get("entidades_exclusivas", 0)])
        ws.append([])

        ws.append([
            "Documento",
            "Substância ativa",
            "Únicas (texto)",
            "Menções (texto)",
            "Únicas (tabelas)",
            "Menções (tabelas)",
            "Menções (total)",
            "% de tabelas",
            "Entidades SÓ em tabelas",
        ])

        for row in linhas:
            total = row["texto_mencoes"] + row["tabelas_mencoes"]
            ws.append([
                row["document"],
                row["active_substance"],
                row["texto_unicas"],
                row["texto_mencoes"],
                row["tabelas_unicas"],
                row["tabelas_mencoes"],
                total,
                self._pct(row["tabelas_mencoes"], total),
                row["exclusivas_de_tabelas_txt"],
            ])

    def _write_guideline_coverage_detail(self, wb: Workbook, result: dict) -> None:
        """Uma linha por par documento×guideline — a folha auditável."""
        ws = wb.create_sheet("Cobertura detalhe")

        ws.append([
            "Documento",
            "Substância ativa",
            "Guideline ID",
            "Fonte",
            "Nome da guideline",
            "Genes da guideline",
            "Genes mencionados no RCM",
            "Genes em falta",
            "Nº esperados",
            "Nº cobertos",
            "% cobertura",
            "Estado",
        ])

        for row in result.get("guideline_coverage", []):
            ws.append([
                row["document"],
                row["active_substance"],
                row["guideline_id"],
                row["source"],
                row["guideline_name"],
                " ; ".join(row["expected_genes"]),
                " ; ".join(row["covered_genes"]),
                " ; ".join(row["missing_genes"]),
                row["n_expected"],
                row["n_covered"],
                self._pct(row["n_covered"], row["n_expected"]),
                row["status"],
            ])

    def _write_guideline_gene_coverage(self, wb: Workbook, result: dict) -> None:
        """Que genes de guideline são mais e menos referidos nos RCMs.

        Duas leituras, lado a lado, porque dão números diferentes e ambas são
        precisas para perguntas diferentes:

          por documento  quantos RCMs deviam mencionar o gene, e quantos o
                         fazem. É a leitura intuitiva e a que responde a
                         "este gene é ignorado pelos rótulos?".
          por posição    o mesmo par fármaco-gene conta uma vez por cada
                         guideline que o cobre. Soma 174 e alimenta a
                         percentagem ponderada do Resumo.

        A diferença não é pequena: o DPYD é esperado em 10 posições mas só em
        2 documentos, porque cinco organismos publicaram uma guideline DPYD
        para a capecitabina e outras cinco para o fluorouracilo.
        """
        ws = wb.create_sheet("Cobertura por gene")

        ws.append([
            "Gene da guideline",
            "Documentos onde é esperado",
            "Documentos que o mencionam",
            "Documentos onde FALTA",
            "% cobertura (por documento)",
            "Posições de guideline",
            "Posições cobertas",
            "% cobertura (ponderada)",
        ])

        for row in result.get("guideline_gene_coverage", []):
            ws.append([
                row["gene"],
                row["documents_expected"],
                row["documents_covered"],
                row["documents_expected"] - row["documents_covered"],
                row["percentage_documents_covered"],
                row["positions_expected"],
                row["positions_covered"],
                row["percentage_positions_covered"],
            ])

    def _write_entidades_clinpgx(self, wb: Workbook, result: dict) -> None:
        ws = wb.create_sheet("Entidades ClinPGx")

        ws.append([
            "Documento",
            "Substância ativa",
            "Entidade",
            "Tipo",
            "Menções reais",
            "Tem informação ClinPGx?",
            "Guideline IDs",
        ])

        for item in result["entidades_clinpgx"]:
            ws.append([
                item.get("document", ""),
                item.get("active_substance", ""),
                item.get("entity", ""),
                item.get("type", ""),
                item.get("total_mentions_real", 0),
                "Sim" if item.get("has_clinpgx") else "Não",
                item.get("guideline_ids", ""),
            ])

    def _write_normalizacoes(self, wb: Workbook, result: dict) -> None:
        ws = wb.create_sheet("Normalizações")

        ws.append([
            "Documento",
            "Substância ativa",
            "Secção",
            "Tipo",
            "Texto original",
            "Texto normalizado",
            "Tipo entidade",
            "Nº menções",
            "Fonte",
        ])

        for item in result["normalizacoes"]:
            ws.append([
                item.get("document", ""),
                item.get("active_substance", ""),
                item.get("section", ""),
                item.get("tipo", ""),
                item.get("original_text", ""),
                item.get("normalized_text", ""),
                item.get("entity_type", ""),
                item.get("total_mentions", ""),
                item.get("source", ""),
            ])

    def _write_genes_agregados(self, wb: Workbook, result: dict) -> None:
        """Genes ordenados por menções, com as dos seus alelos e diplótipos.

        A folha "Genes frequentes" conta só o símbolo escrito por extenso.
        Um gene que os rótulos refiram sobretudo pelo alelo aparece lá em
        baixo: o HLA-B é nomeado directamente num documento e através dos
        seus alelos em 46.

        As menções somam-se porque são registos distintos — o extractor
        anota "CYP2C19" e "CYP2C19*2" separadamente. Os documentos não se
        somam, e por isso a coluna respectiva é a união dos conjuntos: os
        23 documentos que citam CYP2C19*2 citam todos CYP2C19 também, e
        somar daria 298 onde o valor é 275.
        """
        ws = wb.create_sheet("Genes frequentes+")

        ws.append([
            "Gene",
            "Menções do símbolo",
            "Menções via alelo/diplótipo",
            "Menções totais",
            "Nº documentos (união)",
        ])

        for item in result.get("genes_agregados", []) or []:
            ws.append([
                item.get("gene", ""),
                item.get("mencoes_simbolo", 0),
                item.get("mencoes_via_alelo_diplotipo", 0),
                item.get("mencoes_total", 0),
                item.get("documentos_uniao", 0),
            ])

    # ------------------------------------------------------------------
    #        FORMA COMO AS ENTIDADES ESTÃO ESCRITAS
    # ------------------------------------------------------------------
    # O esquema de extração tem genes, alelos, diplótipos e rsIDs. Não tem
    # categoria para genótipo. A consequência aparece aqui: a categoria dos
    # diplótipos acaba ocupada por genótipos (`SLCO1B1 c.521CC`), e a dos
    # alelos por duas grafias do mesmo alelo HLA — a anterior à revisão da
    # nomenclatura, sem dois pontos, e a actual.
    #
    # Estas folhas não corrigem nada. Contam, para que a limitação seja um
    # número em vez de uma impressão, e para que a normalização posterior
    # seja feita sabendo o que está a ser fundido.

    _FORMA_HLA_COM = re.compile(r"^HLA-[A-Z]+\*\d+:\d+$")
    _FORMA_HLA_SEM = re.compile(r"^HLA-[A-Z]+\*\d{4,}$")
    _FORMA_HLA_CURTO = re.compile(r"^HLA-[A-Z]+\*\d{1,2}$")
    _FORMA_STAR = re.compile(r"^[A-Z0-9]+\*\d+[A-Z]?$")
    _FORMA_LETRA = re.compile(r"^[A-Z0-9]+\*[A-Z]$")
    _FORMA_HGVS = re.compile(r"c\.\d+", re.IGNORECASE)
    _FORMA_BASES = re.compile(r"^[ACGT]{2}$")
    _FORMA_GENE_BASES = re.compile(r"^[A-Z0-9]+[ACGT]{2}$")
    _FORMA_DIP = re.compile(r"^[A-Z0-9]+\*\d+[A-Z]?\*\d+[A-Z]?$")
    _FORMA_IMPROV = re.compile(r"^[A-Z0-9]+\*[A-Z]\*[A-Z]$")

    @classmethod
    def _classe_alelo(cls, e: str) -> str:
        if cls._FORMA_HGVS.search(e):     return "variante HGVS, não alelo"
        if cls._FORMA_LETRA.match(e):     return "letra em vez de número (inválido)"
        if cls._FORMA_HLA_COM.match(e):   return "HLA com dois pontos (forma actual)"
        if cls._FORMA_HLA_SEM.match(e):   return "HLA sem dois pontos (forma antiga)"
        if cls._FORMA_HLA_CURTO.match(e): return "HLA incompleto (serótipo)"
        if cls._FORMA_STAR.match(e):      return "alelo estrela válido"
        return "outro"

    @classmethod
    def _classe_diplotipo(cls, e: str) -> str:
        u = e.upper()
        if cls._FORMA_HGVS.search(u):       return "genótipo em HGVS (c.NNN + bases)"
        if cls._FORMA_BASES.match(u):       return "bases soltas, sem gene"
        if cls._FORMA_IMPROV.match(u):      return "improvisação (*base*base)"
        if cls._FORMA_GENE_BASES.match(u):  return "gene + bases (sem c.)"
        if cls._FORMA_DIP.match(u):         return "diplótipo válido"
        return "outro"

    def _write_formas_entidades(self, wb: Workbook, result: dict) -> None:
        grupos = (
            ("alelo", "alelos_frequentes", "star_allele", self._classe_alelo),
            ("diplótipo", "diplotipos_frequentes", "diplotype", self._classe_diplotipo),
        )

        detalhe = wb.create_sheet("Forma das entidades")
        detalhe.append([
            "Categoria", "Entidade", "Classe da forma",
            "Nº documentos", "Nº menções",
        ])

        agregado: list[tuple[str, str, int, int]] = []
        for categoria, chave, campo, classificar in grupos:
            por_classe: dict[str, list[int]] = defaultdict(lambda: [0, 0])
            for item in result.get(chave, []) or []:
                entidade = str(item.get(campo, "") or "")
                if not entidade:
                    continue
                classe = classificar(entidade)
                docs = item.get("documents_count", 0) or 0
                mencoes = item.get("total_mentions", 0) or 0
                detalhe.append([categoria, entidade, classe, docs, mencoes])
                por_classe[classe][0] += 1
                por_classe[classe][1] += mencoes
            for classe, (n_ent, n_men) in por_classe.items():
                agregado.append((categoria, classe, n_ent, n_men))

        resumo = wb.create_sheet("Forma das entidades resumo")
        resumo.append([
            "Categoria", "Classe da forma", "Entidades distintas",
            "Menções", "% das menções da categoria",
        ])
        totais: dict[str, int] = defaultdict(int)
        for categoria, _, _, n_men in agregado:
            totais[categoria] += n_men
        for categoria, classe, n_ent, n_men in sorted(
            agregado, key=lambda r: (r[0], -r[3])
        ):
            total = totais[categoria] or 1
            resumo.append([
                categoria, classe, n_ent, n_men, round(100 * n_men / total, 1),
            ])

    def _write_documentos_excluidos(self, wb: Workbook, result: dict) -> None:
        """Os RCMs que o pipeline não conseguiu ler, e porquê.

        Folha pequena mas que não deve faltar: um documento excluído dos
        denominadores tem de continuar visível, e a razão da exclusão é ela
        própria uma observação sobre o rótulo.
        """
        ws = wb.create_sheet("Documentos excluídos")

        ws.append([
            "Documento",
            "Estado",
            "Causa",
            "Detalhe",
            "Caracteres extraídos",
        ])

        linhas = result["summary"].get("documents_unusable_detail") or []
        for row in linhas:
            ws.append([
                row["documento"],
                row["estado"],
                row["causa"],
                row["detalhe"],
                row["caracteres"],
            ])

        if not linhas:
            ws.append(["(nenhum — todos os documentos foram legíveis)"])

    def _write_genes_sem_tsv(self, wb: Workbook, result: dict) -> None:
        ws = wb.create_sheet("Genes fora do genes.tsv")

        ws.append(["Gene", "Nº documentos", "Documentos"])

        if not result["genes_sem_tsv"]:
            ws.append(["Nenhum gene sem IDs TSV encontrado", 0, ""])
            return

        for item in result["genes_sem_tsv"]:
            ws.append([
                item.get("gene", ""),
                item.get("documents_count", 0),
                " ; ".join(item.get("documents") or []),
            ])

    def _format_workbook(self, wb: Workbook) -> None:
        header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)

        for sheet in wb.worksheets:
            sheet.freeze_panes = "A2"

            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(
                        vertical="top",
                        wrap_text=True,
                    )

            for col_idx in range(1, sheet.max_column + 1):
                col_letter = get_column_letter(col_idx)

                if col_idx == 1:
                    width = 34
                elif col_idx in [2, 3, 4, 5]:
                    width = 30
                else:
                    width = 20

                sheet.column_dimensions[col_letter].width = width


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(
        description="Gera XLS/JSON global de análise PGx a partir de outputs_batch."
    )

    parser.add_argument(
        "outputs_root",
        help="Pasta raiz dos outputs batch. Exemplo: $HOME/Desktop/outputs_batch"
    )

    args = parser.parse_args()

    result = PGxGlobalAnalysis(args.outputs_root).analyse()

    s = result["summary"]

    print("✅ Análise global criada.")
    print(f"Total de documentos processados: {s['total_documents_processed']}")
    print(f"Documentos com Document_Unique: {s['documents_with_document_unique']}")
    if s.get("documents_unusable_excluded"):
        print(f"Excluídos sem texto utilizável: {s['documents_unusable_excluded']}"
              f"  {', '.join(s.get('documents_unusable_names', [])[:5])}")
    print(f"Documentos com PGx: {s['documents_with_pgx']}")
    print(f"Documentos com PGx e entidades: {s['documents_with_pgx_and_entities']}")
    print(f"Documentos com PGx sem entidades: {s['documents_with_pgx_without_entities']}")
    print(f"Total de genes únicos: {s['total_unique_genes']}")
    print(f"Total de menções de genes reais: {s['total_gene_mentions']}")
    print(f"Total de alelos únicos: {s['total_unique_star_alleles']}")
    print(f"Total de menções de alelos reais: {s['total_star_allele_mentions']}")
    print(f"Documentos com guideline ClinPGx: {s['documents_with_guideline']}")
    print(f"Menções de genes com guidelines: {s['total_gene_mentions_with_guideline']}")
    print(f"Menções de entidades com guidelines: {s['total_entity_mentions_with_guideline']}")